import os, sqlite3, hashlib, secrets, csv, io, re
from datetime import datetime, timedelta, date
from pathlib import Path
from typing import Optional, List
import jwt
from fastapi import FastAPI, HTTPException, Depends, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from pydantic import BaseModel, Field
from fastapi.responses import StreamingResponse, Response, FileResponse
from fastapi.staticfiles import StaticFiles
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from openpyxl import Workbook

BASE_DIR=Path(__file__).resolve().parents[2]
DATA_DIR=BASE_DIR/'data'; DATA_DIR.mkdir(exist_ok=True)
DB_PATH=Path(os.getenv('GESTION_DB_PATH',str(DATA_DIR/'gestion_clase_v5_1.db')))
SECRET_KEY=os.getenv('GESTION_SECRET_KEY','CAMBIAR-ESTA-CLAVE-EN-PRODUCCION')
ALGORITHM='HS256'; TOKEN_MINUTES=int(os.getenv('GESTION_TOKEN_MINUTES','480'))
app=FastAPI(title='Sistema de Gestión de Clase V5.1 API',version='1.0.0-etapa20')
app.add_middleware(CORSMiddleware,allow_origins=[x.strip() for x in os.getenv('GESTION_CORS_ORIGINS','http://localhost:8000,http://127.0.0.1:8000').split(',') if x.strip()],allow_credentials=True,allow_methods=['*'],allow_headers=['*'])
bearer=HTTPBearer(auto_error=False)

def conn():
    c=sqlite3.connect(DB_PATH); c.row_factory=sqlite3.Row; c.execute('PRAGMA foreign_keys=ON'); return c

def now(): return datetime.now().isoformat(timespec='seconds')
def hash_password(password):
    salt=secrets.token_hex(16); digest=hashlib.pbkdf2_hmac('sha256',password.encode(),salt.encode(),200000).hex(); return f'pbkdf2$200000${salt}${digest}'
def verify_password(password,stored):
    try:
        _,iterations,salt,digest=stored.split('$',3); candidate=hashlib.pbkdf2_hmac('sha256',password.encode(),salt.encode(),int(iterations)).hex(); return secrets.compare_digest(candidate,digest)
    except Exception:return False

def audit_log(usuario_id, usuario, metodo, ruta, estado_http=None, ip=None, detalle=None):
    try:
        c=conn(); c.execute('INSERT INTO auditoria(usuario_id,usuario,metodo,ruta,estado_http,ip,detalle,fecha) VALUES(?,?,?,?,?,?,?,?)',(usuario_id,usuario,metodo,ruta,estado_http,ip,detalle,now())); c.commit(); c.close()
    except Exception:
        pass

@app.get('/api/health', tags=['sistema'])
def health():
    try:
        c=conn(); r=c.execute('PRAGMA integrity_check').fetchone()[0]; c.close()
        if r != 'ok': raise RuntimeError(r)
        return {'estado':'ok','version':app.version,'base_datos':'ok'}
    except Exception as e:
        raise HTTPException(status_code=503,detail=f'Base de datos no disponible: {e}')

@app.middleware('http')
async def audit_middleware(request:Request, call_next):
    response=None; uid=None; uname=None
    try:
        auth=request.headers.get('authorization','')
        if auth.lower().startswith('bearer '):
            payload=jwt.decode(auth.split(' ',1)[1],SECRET_KEY,algorithms=[ALGORITHM]); uid=int(payload.get('sub'))
            r=rows('SELECT usuario FROM usuarios WHERE id=?',(uid,)); uname=r[0]['usuario'] if r else None
    except Exception:
        pass
    try:
        response=await call_next(request)
        return response
    finally:
        if request.url.path not in ('/api/health','/docs','/openapi.json','/redoc'):
            ip=(request.client.host if request.client else None)
            audit_log(uid,uname,request.method,request.url.path,response.status_code if response else 500,ip)

def init_db():
    c=conn(); cur=c.cursor()
    cur.executescript('''
    CREATE TABLE IF NOT EXISTS usuarios(id INTEGER PRIMARY KEY AUTOINCREMENT,usuario TEXT UNIQUE NOT NULL,nombre TEXT NOT NULL,rol TEXT NOT NULL DEFAULT 'docente',password_hash TEXT NOT NULL,activo INTEGER DEFAULT 1,fecha_registro TEXT NOT NULL,docente_id INTEGER,FOREIGN KEY(docente_id) REFERENCES docentes(id) ON DELETE SET NULL);
    CREATE TABLE IF NOT EXISTS auditoria(id INTEGER PRIMARY KEY AUTOINCREMENT,usuario_id INTEGER,usuario TEXT,metodo TEXT NOT NULL,ruta TEXT NOT NULL,estado_http INTEGER,ip TEXT,detalle TEXT,fecha TEXT NOT NULL,FOREIGN KEY(usuario_id) REFERENCES usuarios(id) ON DELETE SET NULL);
    CREATE INDEX IF NOT EXISTS idx_auditoria_fecha ON auditoria(fecha);
    CREATE INDEX IF NOT EXISTS idx_auditoria_usuario ON auditoria(usuario_id);
    CREATE TABLE IF NOT EXISTS docentes(id INTEGER PRIMARY KEY AUTOINCREMENT,nombre TEXT NOT NULL,cedula TEXT,correo TEXT,telefono TEXT,especialidad TEXT,estado TEXT DEFAULT 'Activo',fecha_registro TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS cursos(id INTEGER PRIMARY KEY AUTOINCREMENT,nombre TEXT NOT NULL,nivel TEXT DEFAULT 'Secundaria',grado TEXT,seccion TEXT,anio_escolar TEXT,docente_id INTEGER,estado TEXT DEFAULT 'Activo',FOREIGN KEY(docente_id) REFERENCES docentes(id) ON DELETE SET NULL);
    CREATE TABLE IF NOT EXISTS asignaturas(id INTEGER PRIMARY KEY AUTOINCREMENT,nombre TEXT NOT NULL,codigo TEXT,grado TEXT,descripcion TEXT,activa INTEGER DEFAULT 1,UNIQUE(nombre,grado));
    CREATE TABLE IF NOT EXISTS estudiantes(id INTEGER PRIMARY KEY AUTOINCREMENT,matricula TEXT UNIQUE NOT NULL,nombre TEXT NOT NULL,curso TEXT,seccion TEXT,correo TEXT,telefono TEXT,estado TEXT DEFAULT 'Activo',fecha_registro TEXT NOT NULL,fecha_nacimiento TEXT,edad_registro INTEGER,curso_id INTEGER,orden_lista INTEGER);
    CREATE TABLE IF NOT EXISTS estudiante_curso(id INTEGER PRIMARY KEY AUTOINCREMENT,estudiante_id INTEGER NOT NULL,curso_id INTEGER NOT NULL,fecha_inicio TEXT NOT NULL,fecha_fin TEXT,UNIQUE(estudiante_id,curso_id),FOREIGN KEY(estudiante_id) REFERENCES estudiantes(id) ON DELETE CASCADE,FOREIGN KEY(curso_id) REFERENCES cursos(id) ON DELETE CASCADE);
    CREATE TABLE IF NOT EXISTS curso_asignatura(id INTEGER PRIMARY KEY AUTOINCREMENT,curso_id INTEGER NOT NULL,asignatura_id INTEGER NOT NULL,docente_id INTEGER,activa INTEGER DEFAULT 1,UNIQUE(curso_id,asignatura_id),FOREIGN KEY(curso_id) REFERENCES cursos(id) ON DELETE CASCADE,FOREIGN KEY(asignatura_id) REFERENCES asignaturas(id) ON DELETE CASCADE,FOREIGN KEY(docente_id) REFERENCES docentes(id) ON DELETE SET NULL);
    CREATE TABLE IF NOT EXISTS competencias(id INTEGER PRIMARY KEY AUTOINCREMENT,asignatura_id INTEGER NOT NULL,codigo TEXT NOT NULL,nombre TEXT NOT NULL,descripcion TEXT NOT NULL,orden INTEGER DEFAULT 1,activa INTEGER DEFAULT 1,UNIQUE(asignatura_id,codigo),FOREIGN KEY(asignatura_id) REFERENCES asignaturas(id) ON DELETE CASCADE);
    CREATE TABLE IF NOT EXISTS competencias_fundamentales(id INTEGER PRIMARY KEY AUTOINCREMENT,codigo TEXT UNIQUE NOT NULL,nombre TEXT NOT NULL,descripcion TEXT NOT NULL,orden INTEGER DEFAULT 1,activa INTEGER DEFAULT 1);
    CREATE TABLE IF NOT EXISTS competencia_fundamental_especifica(id INTEGER PRIMARY KEY AUTOINCREMENT,fundamental_id INTEGER NOT NULL,especifica_id INTEGER NOT NULL,UNIQUE(fundamental_id,especifica_id),FOREIGN KEY(fundamental_id) REFERENCES competencias_fundamentales(id) ON DELETE CASCADE,FOREIGN KEY(especifica_id) REFERENCES competencias(id) ON DELETE CASCADE);
    CREATE TABLE IF NOT EXISTS actividades(id INTEGER PRIMARY KEY AUTOINCREMENT,estudiante_id INTEGER NOT NULL,asignatura_id INTEGER NOT NULL,competencia_id INTEGER NOT NULL,periodo INTEGER NOT NULL,fecha TEXT NOT NULL,titulo TEXT NOT NULL,descripcion TEXT,tipo_evaluacion TEXT NOT NULL,instrumento TEXT NOT NULL,calificacion REAL NOT NULL,observacion TEXT,fecha_registro TEXT NOT NULL,FOREIGN KEY(estudiante_id) REFERENCES estudiantes(id) ON DELETE CASCADE,FOREIGN KEY(asignatura_id) REFERENCES asignaturas(id) ON DELETE CASCADE,FOREIGN KEY(competencia_id) REFERENCES competencias(id) ON DELETE CASCADE);
    CREATE TABLE IF NOT EXISTS calificaciones(id INTEGER PRIMARY KEY AUTOINCREMENT,estudiante_id INTEGER NOT NULL,asignatura_id INTEGER NOT NULL,competencia_id INTEGER NOT NULL,periodo INTEGER NOT NULL,tipo TEXT NOT NULL DEFAULT 'P',calificacion REAL,observacion TEXT,fecha_actualizacion TEXT NOT NULL,UNIQUE(estudiante_id,asignatura_id,competencia_id,periodo,tipo),FOREIGN KEY(estudiante_id) REFERENCES estudiantes(id) ON DELETE CASCADE,FOREIGN KEY(asignatura_id) REFERENCES asignaturas(id) ON DELETE CASCADE,FOREIGN KEY(competencia_id) REFERENCES competencias(id) ON DELETE CASCADE);
    CREATE TABLE IF NOT EXISTS asistencia(id INTEGER PRIMARY KEY AUTOINCREMENT,estudiante_id INTEGER NOT NULL,fecha TEXT NOT NULL,estado TEXT NOT NULL,hora_entrada TEXT,observacion TEXT,UNIQUE(estudiante_id,fecha),FOREIGN KEY(estudiante_id) REFERENCES estudiantes(id) ON DELETE CASCADE);
    CREATE TABLE IF NOT EXISTS incidencias(id INTEGER PRIMARY KEY AUTOINCREMENT,estudiante_id INTEGER NOT NULL,fecha TEXT NOT NULL,hora TEXT,tipo TEXT NOT NULL,gravedad TEXT NOT NULL,descripcion TEXT,accion_tomada TEXT,observacion TEXT,FOREIGN KEY(estudiante_id) REFERENCES estudiantes(id) ON DELETE CASCADE);
    CREATE TABLE IF NOT EXISTS permisos(id INTEGER PRIMARY KEY AUTOINCREMENT,estudiante_id INTEGER NOT NULL,fecha TEXT NOT NULL,tipo TEXT NOT NULL,hora_desde TEXT,hora_hasta TEXT,motivo TEXT,estado TEXT DEFAULT 'Pendiente',observacion TEXT,FOREIGN KEY(estudiante_id) REFERENCES estudiantes(id) ON DELETE CASCADE);
    CREATE TABLE IF NOT EXISTS calendario(id INTEGER PRIMARY KEY AUTOINCREMENT,fecha TEXT NOT NULL,hora_inicio TEXT,hora_fin TEXT,titulo TEXT NOT NULL,tipo TEXT,curso TEXT,seccion TEXT,descripcion TEXT);
    CREATE TABLE IF NOT EXISTS eventos_calendario(id INTEGER PRIMARY KEY AUTOINCREMENT,titulo TEXT NOT NULL,tipo TEXT NOT NULL DEFAULT 'Otro',fecha TEXT NOT NULL,hora_inicio TEXT,hora_fin TEXT,todo_el_dia INTEGER DEFAULT 0,curso_id INTEGER,asignatura_id INTEGER,estudiante_id INTEGER,ubicacion TEXT,descripcion TEXT,recordatorio_minutos INTEGER,estado TEXT DEFAULT 'Programado',creado_por INTEGER,fecha_registro TEXT NOT NULL,FOREIGN KEY(curso_id) REFERENCES cursos(id) ON DELETE CASCADE,FOREIGN KEY(asignatura_id) REFERENCES asignaturas(id) ON DELETE SET NULL,FOREIGN KEY(estudiante_id) REFERENCES estudiantes(id) ON DELETE CASCADE,FOREIGN KEY(creado_por) REFERENCES usuarios(id) ON DELETE SET NULL);
    CREATE INDEX IF NOT EXISTS idx_eventos_calendario_fecha ON eventos_calendario(fecha);
    CREATE INDEX IF NOT EXISTS idx_eventos_calendario_curso ON eventos_calendario(curso_id);
    CREATE TABLE IF NOT EXISTS horario_escolar(id INTEGER PRIMARY KEY AUTOINCREMENT,dia_semana INTEGER NOT NULL,hora_inicio TEXT NOT NULL,hora_fin TEXT NOT NULL,curso_id INTEGER NOT NULL,asignatura_id INTEGER NOT NULL,docente_id INTEGER,aula TEXT,bloque TEXT,activa INTEGER DEFAULT 1,fecha_registro TEXT NOT NULL,FOREIGN KEY(curso_id) REFERENCES cursos(id) ON DELETE CASCADE,FOREIGN KEY(asignatura_id) REFERENCES asignaturas(id) ON DELETE CASCADE,FOREIGN KEY(docente_id) REFERENCES docentes(id) ON DELETE SET NULL);
    CREATE INDEX IF NOT EXISTS idx_horario_curso_dia ON horario_escolar(curso_id,dia_semana,hora_inicio);
    CREATE INDEX IF NOT EXISTS idx_horario_docente_dia ON horario_escolar(docente_id,dia_semana,hora_inicio);
    CREATE TABLE IF NOT EXISTS calificaciones_finales(id INTEGER PRIMARY KEY AUTOINCREMENT,estudiante_id INTEGER NOT NULL,asignatura_id INTEGER NOT NULL,completiva REAL,extraordinaria REAL,especial_cf REAL,especial_ce REAL,situacion TEXT,observacion TEXT,fecha_actualizacion TEXT NOT NULL,UNIQUE(estudiante_id,asignatura_id),FOREIGN KEY(estudiante_id) REFERENCES estudiantes(id) ON DELETE CASCADE,FOREIGN KEY(asignatura_id) REFERENCES asignaturas(id) ON DELETE CASCADE);
    CREATE TABLE IF NOT EXISTS configuracion(clave TEXT PRIMARY KEY,valor TEXT);
    CREATE TABLE IF NOT EXISTS anios_escolares(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nombre TEXT NOT NULL UNIQUE,
        fecha_inicio TEXT,
        fecha_fin TEXT,
        activo INTEGER DEFAULT 0,
        cerrado INTEGER DEFAULT 0,
        fecha_registro TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS periodos_academicos(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        anio_escolar_id INTEGER NOT NULL,
        numero INTEGER NOT NULL,
        nombre TEXT NOT NULL,
        fecha_inicio TEXT,
        fecha_fin TEXT,
        estado TEXT DEFAULT 'Programado',
        fecha_registro TEXT NOT NULL,
        UNIQUE(anio_escolar_id,numero),
        FOREIGN KEY(anio_escolar_id) REFERENCES anios_escolares(id) ON DELETE CASCADE
    );
    CREATE INDEX IF NOT EXISTS idx_periodos_anio ON periodos_academicos(anio_escolar_id,numero);
    CREATE TABLE IF NOT EXISTS planificaciones(
        id INTEGER PRIMARY KEY AUTOINCREMENT, curso_id INTEGER NOT NULL, asignatura_id INTEGER NOT NULL, docente_id INTEGER,
        anio_escolar_id INTEGER NOT NULL, periodo INTEGER NOT NULL, competencia_id INTEGER, titulo TEXT NOT NULL,
        fecha_inicio TEXT, fecha_fin TEXT, tipo_actividad TEXT, instrumento TEXT, descripcion TEXT, recursos TEXT,
        estado TEXT DEFAULT 'Planificada', evento_calendario_id INTEGER, fecha_registro TEXT NOT NULL,
        FOREIGN KEY(curso_id) REFERENCES cursos(id) ON DELETE CASCADE, FOREIGN KEY(asignatura_id) REFERENCES asignaturas(id) ON DELETE CASCADE,
        FOREIGN KEY(docente_id) REFERENCES docentes(id) ON DELETE SET NULL, FOREIGN KEY(anio_escolar_id) REFERENCES anios_escolares(id) ON DELETE CASCADE,
        FOREIGN KEY(competencia_id) REFERENCES competencias(id) ON DELETE SET NULL, FOREIGN KEY(evento_calendario_id) REFERENCES eventos_calendario(id) ON DELETE SET NULL);
    CREATE INDEX IF NOT EXISTS idx_planificaciones_curso_periodo ON planificaciones(curso_id,asignatura_id,periodo);
    CREATE TABLE IF NOT EXISTS estudiantes_datos(id INTEGER PRIMARY KEY AUTOINCREMENT,estudiante_id INTEGER UNIQUE NOT NULL,cedula TEXT,direccion TEXT,municipio TEXT,provincia TEXT,telefono_emergencia TEXT,contacto_emergencia TEXT,parentesco_emergencia TEXT,nacionalidad TEXT,observaciones TEXT,fecha_actualizacion TEXT NOT NULL,FOREIGN KEY(estudiante_id) REFERENCES estudiantes(id) ON DELETE CASCADE);
    CREATE TABLE IF NOT EXISTS familiares(id INTEGER PRIMARY KEY AUTOINCREMENT,estudiante_id INTEGER NOT NULL,nombre TEXT NOT NULL,parentesco TEXT,cedula TEXT,telefono TEXT,correo TEXT,ocupacion TEXT,es_contacto_principal INTEGER DEFAULT 0,observacion TEXT,FOREIGN KEY(estudiante_id) REFERENCES estudiantes(id) ON DELETE CASCADE);
    CREATE TABLE IF NOT EXISTS estudiante_historial(id INTEGER PRIMARY KEY AUTOINCREMENT,estudiante_id INTEGER NOT NULL,curso_id INTEGER,fecha_inicio TEXT NOT NULL,fecha_fin TEXT,motivo TEXT,observacion TEXT,FOREIGN KEY(estudiante_id) REFERENCES estudiantes(id) ON DELETE CASCADE,FOREIGN KEY(curso_id) REFERENCES cursos(id) ON DELETE SET NULL);
    CREATE TABLE IF NOT EXISTS seguimientos(id INTEGER PRIMARY KEY AUTOINCREMENT,estudiante_id INTEGER NOT NULL,fecha TEXT NOT NULL,tipo TEXT NOT NULL,tema TEXT NOT NULL,descripcion TEXT,estado TEXT DEFAULT 'Abierto',prioridad TEXT DEFAULT 'Media',proxima_fecha TEXT,responsable TEXT,fecha_registro TEXT NOT NULL,FOREIGN KEY(estudiante_id) REFERENCES estudiantes(id) ON DELETE CASCADE);
    CREATE TABLE IF NOT EXISTS reuniones_familia(id INTEGER PRIMARY KEY AUTOINCREMENT,estudiante_id INTEGER NOT NULL,fecha TEXT NOT NULL,hora TEXT,tipo TEXT NOT NULL,participantes TEXT,acuerdos TEXT,observaciones TEXT,fecha_registro TEXT NOT NULL,FOREIGN KEY(estudiante_id) REFERENCES estudiantes(id) ON DELETE CASCADE);
    CREATE TABLE IF NOT EXISTS compromisos(id INTEGER PRIMARY KEY AUTOINCREMENT,estudiante_id INTEGER NOT NULL,seguimiento_id INTEGER,descripcion TEXT NOT NULL,fecha_limite TEXT,estado TEXT DEFAULT 'Pendiente',responsable TEXT,fecha_cumplimiento TEXT,observacion TEXT,fecha_registro TEXT NOT NULL,FOREIGN KEY(estudiante_id) REFERENCES estudiantes(id) ON DELETE CASCADE,FOREIGN KEY(seguimiento_id) REFERENCES seguimientos(id) ON DELETE SET NULL);
    CREATE TABLE IF NOT EXISTS alertas_riesgo(id INTEGER PRIMARY KEY AUTOINCREMENT,estudiante_id INTEGER NOT NULL,tipo TEXT NOT NULL,nivel TEXT NOT NULL,titulo TEXT NOT NULL,detalle TEXT,valor REAL,umbral REAL,estado TEXT DEFAULT 'Abierta',fecha_deteccion TEXT NOT NULL,fecha_cierre TEXT,responsable TEXT,observacion TEXT,FOREIGN KEY(estudiante_id) REFERENCES estudiantes(id) ON DELETE CASCADE);
    CREATE TABLE IF NOT EXISTS planes_intervencion(id INTEGER PRIMARY KEY AUTOINCREMENT,estudiante_id INTEGER NOT NULL,alerta_id INTEGER,titulo TEXT NOT NULL,objetivo TEXT,estrategia TEXT,responsable TEXT,fecha_inicio TEXT NOT NULL,fecha_revision TEXT,estado TEXT DEFAULT 'Activo',resultado TEXT,observacion TEXT,fecha_registro TEXT NOT NULL,FOREIGN KEY(estudiante_id) REFERENCES estudiantes(id) ON DELETE CASCADE,FOREIGN KEY(alerta_id) REFERENCES alertas_riesgo(id) ON DELETE SET NULL);
    CREATE TABLE IF NOT EXISTS plan_acciones(id INTEGER PRIMARY KEY AUTOINCREMENT,plan_id INTEGER NOT NULL,descripcion TEXT NOT NULL,responsable TEXT,fecha_limite TEXT,estado TEXT DEFAULT 'Pendiente',evidencia TEXT,fecha_cumplimiento TEXT,observacion TEXT,FOREIGN KEY(plan_id) REFERENCES planes_intervencion(id) ON DELETE CASCADE);
    CREATE TABLE IF NOT EXISTS comunicados(id INTEGER PRIMARY KEY AUTOINCREMENT,titulo TEXT NOT NULL,mensaje TEXT NOT NULL,tipo TEXT DEFAULT 'General',prioridad TEXT DEFAULT 'Normal',autor_usuario_id INTEGER,fecha_publicacion TEXT NOT NULL,fecha_expiracion TEXT,activo INTEGER DEFAULT 1,FOREIGN KEY(autor_usuario_id) REFERENCES usuarios(id) ON DELETE SET NULL);
    CREATE TABLE IF NOT EXISTS comunicado_destinatarios(id INTEGER PRIMARY KEY AUTOINCREMENT,comunicado_id INTEGER NOT NULL,tipo_destinatario TEXT NOT NULL,curso_id INTEGER,estudiante_id INTEGER,UNIQUE(comunicado_id,tipo_destinatario,curso_id,estudiante_id),FOREIGN KEY(comunicado_id) REFERENCES comunicados(id) ON DELETE CASCADE,FOREIGN KEY(curso_id) REFERENCES cursos(id) ON DELETE CASCADE,FOREIGN KEY(estudiante_id) REFERENCES estudiantes(id) ON DELETE CASCADE);
    CREATE TABLE IF NOT EXISTS comunicado_lecturas(id INTEGER PRIMARY KEY AUTOINCREMENT,comunicado_id INTEGER NOT NULL,usuario_id INTEGER NOT NULL,fecha_lectura TEXT NOT NULL,UNIQUE(comunicado_id,usuario_id),FOREIGN KEY(comunicado_id) REFERENCES comunicados(id) ON DELETE CASCADE,FOREIGN KEY(usuario_id) REFERENCES usuarios(id) ON DELETE CASCADE);
    ''')
    cols_u=[r['name'] for r in cur.execute('PRAGMA table_info(usuarios)').fetchall()]
    if 'docente_id' not in cols_u: cur.execute('ALTER TABLE usuarios ADD COLUMN docente_id INTEGER')

    cols=[r['name'] for r in cur.execute('PRAGMA table_info(curso_asignatura)').fetchall()]
    if 'docente_id' not in cols: cur.execute('ALTER TABLE curso_asignatura ADD COLUMN docente_id INTEGER')
    defaults={'anio_escolar':'2026-2027','max_calificacion':'100','rp_reemplaza_p':'1','final_promedio_competencias':'1','nombre_centro':'Centro Educativo','umbral_aprobacion':'70','peso_completiva_cf':'50','peso_completiva_cec':'50','peso_extra_cf':'30','peso_extra_ceex':'70'}
    for k,v in defaults.items():cur.execute('INSERT OR IGNORE INTO configuracion VALUES (?,?)',(k,v))
    anio_cfg=cur.execute("SELECT valor FROM configuracion WHERE clave='anio_escolar'").fetchone()[0]
    cur.execute('INSERT OR IGNORE INTO anios_escolares(nombre,activo,fecha_registro) VALUES(?,?,?)',(anio_cfg,1,now()))
    active_id=cur.execute("SELECT id FROM anios_escolares WHERE nombre=?",(anio_cfg,)).fetchone()[0]
    if cur.execute('SELECT COUNT(*) FROM anios_escolares WHERE activo=1').fetchone()[0]==0:
        cur.execute('UPDATE anios_escolares SET activo=1 WHERE id=?',(active_id,))
    for n in range(1,5):
        cur.execute('INSERT OR IGNORE INTO periodos_academicos(anio_escolar_id,numero,nombre,estado,fecha_registro) VALUES(?,?,?,?,?)',(active_id,n,f'Período {n}','Programado',now()))
    if cur.execute('SELECT COUNT(*) FROM usuarios').fetchone()[0]==0:cur.execute('INSERT INTO usuarios(usuario,nombre,rol,password_hash,fecha_registro) VALUES(?,?,?,?,?)',('admin','Administrador','admin',hash_password('admin123'),now()))
    c.commit(); c.close()
init_db()

def rows(sql,args=()):
    c=conn(); out=[dict(r) for r in c.execute(sql,args).fetchall()]; c.close(); return out

def public_user(r): return {'id':r['id'],'usuario':r['usuario'],'nombre':r['nombre'],'rol':r['rol'],'activo':bool(r['activo']),'fecha_registro':r['fecha_registro'],'docente_id':r['docente_id'] if 'docente_id' in r.keys() else None}
def current_user(creds:HTTPAuthorizationCredentials=Depends(bearer)):
    if not creds: raise HTTPException(401,'Token requerido')
    try: payload=jwt.decode(creds.credentials,SECRET_KEY,algorithms=[ALGORITHM]); uid=int(payload['sub'])
    except Exception: raise HTTPException(401,'Token inválido o expirado')
    c=conn(); r=c.execute('SELECT * FROM usuarios WHERE id=?',(uid,)).fetchone(); c.close()
    if not r or not r['activo']: raise HTTPException(401,'Usuario inactivo o inexistente')
    return public_user(r)
def admin_user(u=Depends(current_user)):
    if u['rol']!='admin': raise HTTPException(403,'Solo un administrador puede realizar esta operación')
    return u

class LoginIn(BaseModel): usuario:str=Field(min_length=1,max_length=60); password:str=Field(min_length=1,max_length=200)
class DocenteIn(BaseModel): nombre:str=Field(min_length=2,max_length=150); cedula:Optional[str]=None; correo:Optional[str]=None; telefono:Optional[str]=None; especialidad:Optional[str]=None; estado:str='Activo'
class CursoIn(BaseModel): nombre:str; nivel:str='Secundaria'; grado:Optional[str]=None; seccion:Optional[str]=None; anio_escolar:Optional[str]=None; docente_id:Optional[int]=None
class AsignaturaIn(BaseModel): nombre:str; codigo:Optional[str]=None; grado:Optional[str]=None; descripcion:Optional[str]=None
class LinkIn(BaseModel): curso_id:int; asignatura_id:int; docente_id:Optional[int]=None
class EstudianteIn(BaseModel): nombres:str; apellidos:str; fecha_nacimiento:date; curso_id:int; correo:Optional[str]=None; telefono:Optional[str]=None; estado:str='Activo'
class EstudianteDatosIn(BaseModel): cedula:Optional[str]=None; direccion:Optional[str]=None; municipio:Optional[str]=None; provincia:Optional[str]=None; telefono_emergencia:Optional[str]=None; contacto_emergencia:Optional[str]=None; parentesco_emergencia:Optional[str]=None; nacionalidad:Optional[str]=None; observaciones:Optional[str]=None
class FamiliarIn(BaseModel): nombre:str; parentesco:Optional[str]=None; cedula:Optional[str]=None; telefono:Optional[str]=None; correo:Optional[str]=None; ocupacion:Optional[str]=None; es_contacto_principal:bool=False; observacion:Optional[str]=None
class TrasladoIn(BaseModel): nuevo_curso_id:int; fecha:date; motivo:Optional[str]=None; observacion:Optional[str]=None
class SeguimientoIn(BaseModel): estudiante_id:int; fecha:date; tipo:str; tema:str; descripcion:Optional[str]=None; estado:str='Abierto'; prioridad:str='Media'; proxima_fecha:Optional[date]=None; responsable:Optional[str]=None
class ReunionIn(BaseModel): estudiante_id:int; fecha:date; hora:Optional[str]=None; tipo:str='Reunión con familia'; participantes:Optional[str]=None; acuerdos:Optional[str]=None; observaciones:Optional[str]=None
class CompromisoIn(BaseModel): estudiante_id:int; seguimiento_id:Optional[int]=None; descripcion:str; fecha_limite:Optional[date]=None; estado:str='Pendiente'; responsable:Optional[str]=None; fecha_cumplimiento:Optional[date]=None; observacion:Optional[str]=None
class AlertaEstadoIn(BaseModel):
    estado:str
    responsable:Optional[str]=None
    observacion:Optional[str]=None
class PlanIn(BaseModel):
    estudiante_id:int; alerta_id:Optional[int]=None; titulo:str; objetivo:Optional[str]=None; estrategia:Optional[str]=None; responsable:Optional[str]=None; fecha_inicio:date; fecha_revision:Optional[date]=None; estado:str='Activo'; resultado:Optional[str]=None; observacion:Optional[str]=None
class PlanAccionIn(BaseModel):
    plan_id:int; descripcion:str; responsable:Optional[str]=None; fecha_limite:Optional[date]=None; estado:str='Pendiente'; evidencia:Optional[str]=None; fecha_cumplimiento:Optional[date]=None; observacion:Optional[str]=None
class ComunicadoDestinatarioIn(BaseModel):
    tipo_destinatario:str='todos'; curso_id:Optional[int]=None; estudiante_id:Optional[int]=None
class ComunicadoIn(BaseModel):
    titulo:str=Field(min_length=2,max_length=200); mensaje:str=Field(min_length=1,max_length=10000); tipo:str='General'; prioridad:str='Normal'; fecha_expiracion:Optional[date]=None; activo:bool=True; destinatarios:List[ComunicadoDestinatarioIn]=Field(default_factory=list)
class PlanificacionIn(BaseModel):
    curso_id:int
    asignatura_id:int
    docente_id:Optional[int]=None
    anio_escolar_id:int
    periodo:int=Field(ge=1,le=4)
    competencia_id:Optional[int]=None
    titulo:str
    fecha_inicio:Optional[date]=None
    fecha_fin:Optional[date]=None
    tipo_actividad:Optional[str]=None
    instrumento:Optional[str]=None
    descripcion:Optional[str]=None
    recursos:Optional[str]=None
    estado:str='Planificada'
    crear_evento:bool=False

class HorarioIn(BaseModel): dia_semana:int=Field(ge=1,le=7); hora_inicio:str; hora_fin:str; curso_id:int; asignatura_id:int; docente_id:Optional[int]=None; aula:Optional[str]=None; bloque:Optional[str]=None; activa:bool=True
class EventoCalendarioIn(BaseModel):
    titulo:str=Field(min_length=2,max_length=200); tipo:str='Otro'; fecha:date; hora_inicio:Optional[str]=None; hora_fin:Optional[str]=None; todo_el_dia:bool=False; curso_id:Optional[int]=None; asignatura_id:Optional[int]=None; estudiante_id:Optional[int]=None; ubicacion:Optional[str]=None; descripcion:Optional[str]=None; recordatorio_minutos:Optional[int]=Field(default=None,ge=0,le=10080); estado:str='Programado'

class CompetenciaIn(BaseModel): asignatura_id:int; codigo:str; nombre:str; descripcion:str=''; orden:int=1; activa:bool=True
class FundamentalIn(BaseModel): codigo:str; nombre:str; descripcion:str=''; orden:int=1; activa:bool=True
class MapeoIn(BaseModel): fundamental_id:int; especifica_id:int
class ActividadIn(BaseModel): estudiante_id:int; asignatura_id:int; competencia_id:int; periodo:int=Field(ge=1,le=4); fecha:date; titulo:str; descripcion:Optional[str]=None; tipo_evaluacion:str; instrumento:str; calificacion:float=Field(ge=0,le=100); observacion:Optional[str]=None
class CalificacionIn(BaseModel): estudiante_id:int; asignatura_id:int; competencia_id:int; periodo:int=Field(ge=1,le=4); tipo:str='P'; calificacion:Optional[float]=Field(default=None,ge=0,le=100); observacion:Optional[str]=None
class CalificacionLoteItem(BaseModel): estudiante_id:int; competencia_id:int; calificacion:Optional[float]=Field(default=None,ge=0,le=100); observacion:Optional[str]=None
class CalificacionLoteIn(BaseModel): curso_id:int; asignatura_id:int; periodo:int=Field(ge=1,le=4); tipo:str='P'; registros:List[CalificacionLoteItem]
class ActividadUpdateIn(BaseModel): fecha:date; titulo:str; descripcion:Optional[str]=None; tipo_evaluacion:str; instrumento:str; calificacion:float=Field(ge=0,le=100); observacion:Optional[str]=None
class AsistenciaIn(BaseModel): estudiante_id:int; fecha:date; estado:str; hora_entrada:Optional[str]=None; observacion:Optional[str]=None
class AsistenciaLoteIn(BaseModel): curso_id:int; fecha:date; registros:List[dict]
class IncidenciaIn(BaseModel): estudiante_id:int; fecha:date; hora:Optional[str]=None; tipo:str; gravedad:str='Leve'; descripcion:Optional[str]=None; accion_tomada:Optional[str]=None; observacion:Optional[str]=None
class PermisoIn(BaseModel): estudiante_id:int; fecha:date; tipo:str; hora_desde:Optional[str]=None; hora_hasta:Optional[str]=None; motivo:Optional[str]=None; estado:str='Pendiente'; observacion:Optional[str]=None
class PermisoUpdateIn(BaseModel): fecha:Optional[date]=None; tipo:Optional[str]=None; hora_desde:Optional[str]=None; hora_hasta:Optional[str]=None; motivo:Optional[str]=None; estado:Optional[str]=None; observacion:Optional[str]=None
class CalificacionFinalIn(BaseModel):
    estudiante_id:int
    asignatura_id:int
    completiva:Optional[float]=Field(default=None,ge=0,le=100)
    extraordinaria:Optional[float]=Field(default=None,ge=0,le=100)
    especial_cf:Optional[float]=Field(default=None,ge=0,le=100)
    especial_ce:Optional[float]=Field(default=None,ge=0,le=100)
    situacion:Optional[str]=None
    observacion:Optional[str]=None
class ConfigUpdate(BaseModel):
    valor:str
class AnioEscolarIn(BaseModel):
    nombre:str=Field(min_length=4,max_length=40)
    fecha_inicio:Optional[date]=None
    fecha_fin:Optional[date]=None
    activo:bool=False
    cerrado:bool=False
class PeriodoAcademicoIn(BaseModel):
    anio_escolar_id:int
    numero:int=Field(ge=1,le=4)
    nombre:str=Field(min_length=2,max_length=80)
    fecha_inicio:Optional[date]=None
    fecha_fin:Optional[date]=None
    estado:str='Programado'

class UsuarioCreate(BaseModel):
    usuario:str=Field(min_length=3,max_length=60); nombre:str=Field(min_length=2,max_length=150); rol:str='docente'; password:str=Field(min_length=6,max_length=200)
class UsuarioUpdate(BaseModel):
    nombre:str=Field(min_length=2,max_length=150); rol:str='docente'; activo:bool=True
class PasswordUpdate(BaseModel):
    password:str=Field(min_length=8,max_length=200)
class VincularDocenteUsuarioIn(BaseModel):
    docente_id:Optional[int]=None
class InstitucionUpdate(BaseModel):
    nombre_centro:str=Field(min_length=2,max_length=200)
    codigo_centro:Optional[str]=None
    distrito:Optional[str]=None
    regional:Optional[str]=None
    direccion:Optional[str]=None
    telefono:Optional[str]=None
    correo:Optional[str]=None
    director:Optional[str]=None

class ImportarDocentesIn(BaseModel):
    csv_text:str
class ImportarEstudiantesIn(BaseModel):
    csv_text:str

@app.get('/api/health', tags=['sistema'])
def health():
    try:
        c=conn(); integrity=c.execute('PRAGMA integrity_check').fetchone()[0]; c.close()
        if integrity != 'ok': raise RuntimeError(integrity)
        return {'estado':'ok','version':app.version,'base_datos':'ok'}
    except Exception as e:
        raise HTTPException(status_code=503,detail=f'Base de datos no disponible: {e}')
@app.post('/api/auth/login')
def login(x:LoginIn):
    c=conn();r=c.execute('SELECT * FROM usuarios WHERE usuario=? AND activo=1',(x.usuario.strip(),)).fetchone();c.close()
    if not r or not verify_password(x.password,r['password_hash']):raise HTTPException(401,'Usuario o contraseña incorrectos')
    token=jwt.encode({'sub':str(r['id']),'exp':datetime.utcnow()+timedelta(minutes=TOKEN_MINUTES)},SECRET_KEY,algorithm=ALGORITHM)
    return {'access_token':token,'token_type':'bearer',**public_user(r)}
@app.get('/api/auth/me')
def me(u=Depends(current_user)):return u

@app.get('/api/docentes')
def docentes(u=Depends(current_user)):
    return rows("SELECT id,nombre,cedula,correo,telefono,especialidad,estado,fecha_registro FROM docentes ORDER BY nombre COLLATE NOCASE")
@app.post('/api/docentes')
def crear_docente(x:DocenteIn,u=Depends(admin_user)):
    c=conn();cur=c.execute('INSERT INTO docentes(nombre,cedula,correo,telefono,especialidad,estado,fecha_registro) VALUES(?,?,?,?,?,?,?)',(x.nombre.strip(),x.cedula,x.correo,x.telefono,x.especialidad,x.estado,now()));c.commit();r=c.execute('SELECT * FROM docentes WHERE id=?',(cur.lastrowid,)).fetchone();c.close();return dict(r)
@app.put('/api/docentes/{docente_id}')
def actualizar_docente(docente_id:int,x:DocenteIn,u=Depends(admin_user)):
    c=conn();r=c.execute('SELECT id FROM docentes WHERE id=?',(docente_id,)).fetchone()
    if not r:c.close();raise HTTPException(404,'Docente no encontrado')
    c.execute('UPDATE docentes SET nombre=?,cedula=?,correo=?,telefono=?,especialidad=?,estado=? WHERE id=?',(x.nombre.strip(),x.cedula,x.correo,x.telefono,x.especialidad,x.estado,docente_id));c.commit();out=dict(c.execute('SELECT * FROM docentes WHERE id=?',(docente_id,)).fetchone());c.close();return out
@app.delete('/api/docentes/{docente_id}')
def desactivar_docente(docente_id:int,u=Depends(admin_user)):
    c=conn();r=c.execute('SELECT id FROM docentes WHERE id=?',(docente_id,)).fetchone()
    if not r:c.close();raise HTTPException(404,'Docente no encontrado')
    c.execute("UPDATE docentes SET estado='Inactivo' WHERE id=?",(docente_id,));c.commit();c.close();return {'ok':True}

def docente_scope(u):
    if u['rol']=='admin': return None
    if not u.get('docente_id'): return {'docente_id':None}
    return {'docente_id':u['docente_id']}

def allowed_course_ids(u):
    if u['rol']=='admin': return None
    did=u.get('docente_id')
    if not did: return []
    rs=rows('SELECT id FROM cursos WHERE docente_id=? UNION SELECT curso_id AS id FROM curso_asignatura WHERE docente_id=? AND activa=1',(did,did))
    return [r['id'] for r in rs]

def allowed_subject_pairs(u):
    if u['rol']=='admin': return None
    did=u.get('docente_id')
    if not did: return []
    return rows('SELECT curso_id,asignatura_id FROM curso_asignatura WHERE docente_id=? AND activa=1',(did,))

def can_access_student(u,student_id):
    if u['rol']=='admin': return True
    return bool(rows('''SELECT e.id FROM estudiantes e LEFT JOIN cursos c ON c.id=e.curso_id LEFT JOIN curso_asignatura ca ON ca.curso_id=e.curso_id AND ca.activa=1 WHERE e.id=? AND (c.docente_id=? OR ca.docente_id=?)''',(student_id,u.get('docente_id'),u.get('docente_id'))))

@app.get('/api/cursos')
def cursos(u=Depends(current_user)):
    sql="SELECT c.*,d.nombre docente_nombre,d.especialidad docente_especialidad FROM cursos c LEFT JOIN docentes d ON d.id=c.docente_id"
    ids=allowed_course_ids(u)
    if ids is not None:
        if not ids:return []
        sql+=' WHERE c.id IN ('+','.join('?'*len(ids))+')'; sql+=' ORDER BY c.nombre COLLATE NOCASE'; return rows(sql,ids)
    sql+=' ORDER BY c.nombre COLLATE NOCASE'; return rows(sql)
@app.post('/api/cursos')
def crear_curso(x:CursoIn,u=Depends(current_user)):
    if x.docente_id:
        if not rows("SELECT id FROM docentes WHERE id=? AND estado='Activo'",(x.docente_id,)):raise HTTPException(400,'El docente/tutor no está activo o no existe')
    c=conn();cur=c.execute('INSERT INTO cursos(nombre,nivel,grado,seccion,anio_escolar,docente_id) VALUES(?,?,?,?,?,?)',(x.nombre,x.nivel,x.grado,x.seccion,x.anio_escolar,x.docente_id));c.commit();r=dict(c.execute('SELECT * FROM cursos WHERE id=?',(cur.lastrowid,)).fetchone());c.close();return r

@app.get('/api/asignaturas')
def asignaturas(curso_id:Optional[int]=None,u=Depends(current_user)):
    if curso_id:
        if u['rol']!='admin' and curso_id not in (allowed_course_ids(u) or []): return []
        return rows("SELECT a.id,a.nombre,a.codigo,a.grado,a.descripcion,ca.docente_id,d.nombre docente_nombre,d.especialidad docente_especialidad,ca.activa FROM asignaturas a JOIN curso_asignatura ca ON ca.asignatura_id=a.id LEFT JOIN docentes d ON d.id=ca.docente_id WHERE ca.curso_id=? AND ca.activa=1 ORDER BY a.nombre COLLATE NOCASE",(curso_id,))
    if u['rol']=='admin': return rows('SELECT * FROM asignaturas WHERE activa=1 ORDER BY grado,nombre COLLATE NOCASE')
    pairs=allowed_subject_pairs(u)
    if not pairs:return []
    ids=sorted(set(p['asignatura_id'] for p in pairs)); return rows('SELECT * FROM asignaturas WHERE activa=1 AND id IN ('+','.join('?'*len(ids))+') ORDER BY grado,nombre COLLATE NOCASE',ids)
@app.post('/api/asignaturas')
def crear_asignatura(x:AsignaturaIn,u=Depends(current_user)):
    c=conn()
    try:cur=c.execute('INSERT INTO asignaturas(nombre,codigo,grado,descripcion) VALUES(?,?,?,?)',(x.nombre,x.codigo,x.grado,x.descripcion));c.commit()
    except sqlite3.IntegrityError:c.close();raise HTTPException(409,'La asignatura ya existe para ese grado')
    r=dict(c.execute('SELECT * FROM asignaturas WHERE id=?',(cur.lastrowid,)).fetchone());c.close();return r
@app.post('/api/cursos/asignaturas')
def vincular(x:LinkIn,u=Depends(current_user)):
    if not rows('SELECT id FROM cursos WHERE id=?',(x.curso_id,)):raise HTTPException(404,'Clase no encontrada')
    if not rows('SELECT id FROM asignaturas WHERE id=? AND activa=1',(x.asignatura_id,)):raise HTTPException(404,'Asignatura no encontrada o inactiva')
    if x.docente_id and not rows("SELECT id FROM docentes WHERE id=? AND estado='Activo'",(x.docente_id,)):raise HTTPException(400,'El docente seleccionado no está activo o no existe')
    c=conn();c.execute('INSERT OR IGNORE INTO curso_asignatura(curso_id,asignatura_id,docente_id,activa) VALUES(?,?,?,1)',(x.curso_id,x.asignatura_id,x.docente_id));c.execute('UPDATE curso_asignatura SET docente_id=?,activa=1 WHERE curso_id=? AND asignatura_id=?',(x.docente_id,x.curso_id,x.asignatura_id));c.commit();c.close();return {'ok':True,'message':'Asignatura asociada con docente responsable'}

@app.get('/api/estudiantes')
def estudiantes(curso_id:Optional[int]=None,u=Depends(current_user)):
    sql='SELECT e.*,c.nombre curso_nombre,c.grado,c.seccion FROM estudiantes e LEFT JOIN cursos c ON c.id=e.curso_id';args=[]
    if u['rol']=='admin':
        if curso_id:sql+=' WHERE e.curso_id=?';args.append(curso_id)
    else:
        ids=allowed_course_ids(u) or []
        if not ids:return []
        sql+=' WHERE e.curso_id IN ('+','.join('?'*len(ids))+')';args.extend(ids)
        if curso_id:
            if curso_id not in ids:return []
            sql+=' AND e.curso_id=?';args.append(curso_id)
    sql+=' ORDER BY e.nombre COLLATE NOCASE';return rows(sql,args)
def age(d):
    t=date.today();return t.year-d.year-((t.month,t.day)<(d.month,d.day))
def initial(s):return next((w[0].upper() for w in s.strip().split() if w),'X')
@app.post('/api/estudiantes')
def crear_estudiante(x:EstudianteIn,u=Depends(current_user)):
    c=conn();cur=c.cursor();course=cur.execute('SELECT * FROM cursos WHERE id=?',(x.curso_id,)).fetchone()
    if not course:c.close();raise HTTPException(404,'Clase no encontrada')
    n=cur.execute('SELECT COALESCE(MAX(orden_lista),0)+1 FROM estudiantes WHERE curso_id=?',(x.curso_id,)).fetchone()[0]
    base=f'{initial(x.nombres)}{initial(x.apellidos)}{x.fecha_nacimiento.year}-{n}';mat=base;k=2
    while cur.execute('SELECT 1 FROM estudiantes WHERE matricula=?',(mat,)).fetchone():mat=f'{base}-{k}';k+=1
    nombre=f'{x.nombres.strip()} {x.apellidos.strip()}';a=age(x.fecha_nacimiento)
    cur.execute('INSERT INTO estudiantes(matricula,nombre,curso,seccion,correo,telefono,estado,fecha_registro,fecha_nacimiento,edad_registro,curso_id,orden_lista) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',(mat,nombre,course['grado'],course['seccion'],x.correo,x.telefono,x.estado,now(),x.fecha_nacimiento.isoformat(),a,x.curso_id,n));sid=cur.lastrowid
    cur.execute('INSERT OR IGNORE INTO estudiante_curso(estudiante_id,curso_id,fecha_inicio) VALUES(?,?,?)',(sid,x.curso_id,date.today().isoformat()));c.commit();c.close();return {'id':sid,'matricula':mat,'nombre':nombre,'edad_registro':a,'curso_id':x.curso_id,'orden_lista':n}
@app.get('/api/estudiantes/{student_id}/expediente')
def expediente_estudiante(student_id:int,u=Depends(current_user)):
    if not can_access_student(u,student_id): raise HTTPException(403,'No tiene acceso a este estudiante')
    st=rows('''SELECT e.*,c.nombre curso_nombre,c.grado,c.seccion,c.anio_escolar FROM estudiantes e LEFT JOIN cursos c ON c.id=e.curso_id WHERE e.id=?''',(student_id,))
    if not st: raise HTTPException(404,'Estudiante no encontrado')
    datos=rows('SELECT * FROM estudiantes_datos WHERE estudiante_id=?',(student_id,))
    fam=rows('SELECT * FROM familiares WHERE estudiante_id=? ORDER BY es_contacto_principal DESC,nombre COLLATE NOCASE',(student_id,))
    hist=rows('''SELECT h.id,h.estudiante_id,h.curso_id,h.fecha_inicio,h.fecha_fin,NULL motivo,NULL observacion,c.nombre curso_nombre,c.grado,c.seccion FROM estudiante_curso h LEFT JOIN cursos c ON c.id=h.curso_id WHERE h.estudiante_id=? ORDER BY h.fecha_inicio DESC''',(student_id,))
    asist=rows('''SELECT estado,COUNT(*) cantidad FROM asistencia WHERE estudiante_id=? GROUP BY estado ORDER BY estado''',(student_id,))
    inc=rows('SELECT id,fecha,tipo,gravedad,descripcion,accion_tomada,observacion FROM incidencias WHERE estudiante_id=? ORDER BY fecha DESC,id DESC LIMIT 50',(student_id,))
    per=rows('SELECT id,fecha,tipo,hora_desde,hora_hasta,motivo,estado,observacion FROM permisos WHERE estudiante_id=? ORDER BY fecha DESC,id DESC LIMIT 50',(student_id,))
    return {'estudiante':st[0],'datos':datos[0] if datos else None,'familiares':fam,'historial':hist,'asistencia':asist,'incidencias':inc,'permisos':per}

@app.put('/api/estudiantes/{student_id}/datos')
def actualizar_datos_estudiante(student_id:int,x:EstudianteDatosIn,u=Depends(current_user)):
    if not can_access_student(u,student_id): raise HTTPException(403,'No tiene acceso a este estudiante')
    c=conn(); c.execute('''INSERT INTO estudiantes_datos(estudiante_id,cedula,direccion,municipio,provincia,telefono_emergencia,contacto_emergencia,parentesco_emergencia,nacionalidad,observaciones,fecha_actualizacion) VALUES(?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(estudiante_id) DO UPDATE SET cedula=excluded.cedula,direccion=excluded.direccion,municipio=excluded.municipio,provincia=excluded.provincia,telefono_emergencia=excluded.telefono_emergencia,contacto_emergencia=excluded.contacto_emergencia,parentesco_emergencia=excluded.parentesco_emergencia,nacionalidad=excluded.nacionalidad,observaciones=excluded.observaciones,fecha_actualizacion=excluded.fecha_actualizacion''',(student_id,x.cedula,x.direccion,x.municipio,x.provincia,x.telefono_emergencia,x.contacto_emergencia,x.parentesco_emergencia,x.nacionalidad,x.observaciones,now())); c.commit(); out=dict(c.execute('SELECT * FROM estudiantes_datos WHERE estudiante_id=?',(student_id,)).fetchone()); c.close(); return out

@app.post('/api/estudiantes/{student_id}/familiares')
def agregar_familiar(student_id:int,x:FamiliarIn,u=Depends(current_user)):
    if not can_access_student(u,student_id): raise HTTPException(403,'No tiene acceso a este estudiante')
    c=conn();
    if x.es_contacto_principal: c.execute('UPDATE familiares SET es_contacto_principal=0 WHERE estudiante_id=?',(student_id,))
    cur=c.execute('INSERT INTO familiares(estudiante_id,nombre,parentesco,cedula,telefono,correo,ocupacion,es_contacto_principal,observacion) VALUES(?,?,?,?,?,?,?,?,?)',(student_id,x.nombre.strip(),x.parentesco,x.cedula,x.telefono,x.correo,x.ocupacion,int(x.es_contacto_principal),x.observacion)); c.commit(); r=dict(c.execute('SELECT * FROM familiares WHERE id=?',(cur.lastrowid,)).fetchone()); c.close(); return r

@app.delete('/api/estudiantes/{student_id}/familiares/{familiar_id}')
def eliminar_familiar(student_id:int,familiar_id:int,u=Depends(current_user)):
    if not can_access_student(u,student_id): raise HTTPException(403,'No tiene acceso a este estudiante')
    c=conn(); r=c.execute('SELECT id FROM familiares WHERE id=? AND estudiante_id=?',(familiar_id,student_id)).fetchone()
    if not r: c.close(); raise HTTPException(404,'Familiar no encontrado')
    c.execute('DELETE FROM familiares WHERE id=?',(familiar_id,)); c.commit(); c.close(); return {'ok':True}

@app.post('/api/estudiantes/{student_id}/traslado')
def trasladar_estudiante(student_id:int,x:TrasladoIn,u=Depends(current_user)):
    if not can_access_student(u,student_id): raise HTTPException(403,'No tiene acceso al estudiante')
    if not rows('SELECT id FROM cursos WHERE id=?',(x.nuevo_curso_id,)): raise HTTPException(404,'Nueva clase no encontrada')
    if u['rol']!='admin' and x.nuevo_curso_id not in (allowed_course_ids(u) or []): raise HTTPException(403,'No tiene acceso a la nueva clase')
    c=conn(); old=c.execute('SELECT curso_id FROM estudiantes WHERE id=?',(student_id,)).fetchone()
    if not old: c.close(); raise HTTPException(404,'Estudiante no encontrado')
    old_id=old['curso_id']; c.execute('UPDATE estudiante_curso SET fecha_fin=? WHERE estudiante_id=? AND curso_id=? AND fecha_fin IS NULL',(x.fecha.isoformat(),student_id,old_id))
    c.execute('INSERT OR IGNORE INTO estudiante_curso(estudiante_id,curso_id,fecha_inicio,fecha_fin) VALUES(?,?,?,NULL)',(student_id,x.nuevo_curso_id,x.fecha.isoformat()))
    c.execute('INSERT INTO estudiante_historial(estudiante_id,curso_id,fecha_inicio,fecha_fin,motivo,observacion) VALUES(?,?,?,?,?,?)',(student_id,old_id,x.fecha.isoformat(),x.fecha.isoformat(),x.motivo,x.observacion))
    c.execute('UPDATE estudiantes SET curso_id=?,curso=(SELECT grado FROM cursos WHERE id=?),seccion=(SELECT seccion FROM cursos WHERE id=?),orden_lista=(SELECT COALESCE(MAX(orden_lista),0)+1 FROM estudiantes WHERE curso_id=?) WHERE id=?',(x.nuevo_curso_id,x.nuevo_curso_id,x.nuevo_curso_id,x.nuevo_curso_id,student_id)); c.commit(); out=dict(c.execute('SELECT e.*,c.nombre curso_nombre,c.grado,c.seccion FROM estudiantes e LEFT JOIN cursos c ON c.id=e.curso_id WHERE e.id=?',(student_id,)).fetchone()); c.close(); return out

@app.get('/api/seguimientos')
def get_seguimientos(student_id:Optional[int]=None,u=Depends(current_user)):
    if student_id is not None and not can_access_student(u,student_id): raise HTTPException(403,'Sin acceso al estudiante')
    if u['rol']=='admin':
        q='''SELECT sg.*,e.matricula,e.nombre estudiante FROM seguimientos sg JOIN estudiantes e ON e.id=sg.estudiante_id WHERE 1=1'''; args=[]
    else:
        ids=allowed_course_ids(u) or []
        if not ids: return []
        ph=','.join('?'*len(ids)); q=f'''SELECT sg.*,e.matricula,e.nombre estudiante FROM seguimientos sg JOIN estudiantes e ON e.id=sg.estudiante_id WHERE e.curso_id IN ({ph})'''; args=ids
    if student_id is not None: q+=' AND sg.estudiante_id=?'; args.append(student_id)
    q+=' ORDER BY sg.fecha DESC,sg.id DESC'
    return rows(q,args)

@app.post('/api/seguimientos')
def post_seguimiento(x:SeguimientoIn,u=Depends(current_user)):
    if not can_access_student(u,x.estudiante_id): raise HTTPException(403,'Sin acceso al estudiante')
    cur=conn(); cur.execute('INSERT INTO seguimientos(estudiante_id,fecha,tipo,tema,descripcion,estado,prioridad,proxima_fecha,responsable,fecha_registro) VALUES(?,?,?,?,?,?,?,?,?,?)',(x.estudiante_id,x.fecha.isoformat(),x.tipo,x.tema,x.descripcion,x.estado,x.prioridad,x.proxima_fecha.isoformat() if x.proxima_fecha else None,x.responsable,now())); cur.commit(); return {'id':cur.lastrowid}

@app.put('/api/seguimientos/{sid}')
def put_seguimiento(sid:int,x:SeguimientoIn,u=Depends(current_user)):
    r=rows('SELECT estudiante_id FROM seguimientos WHERE id=?',(sid,));
    if not r: raise HTTPException(404,'Seguimiento no encontrado')
    if not can_access_student(u,r[0]['estudiante_id']): raise HTTPException(403,'Sin acceso')
    cur=conn(); cur.execute('UPDATE seguimientos SET fecha=?,tipo=?,tema=?,descripcion=?,estado=?,prioridad=?,proxima_fecha=?,responsable=? WHERE id=?',(x.fecha.isoformat(),x.tipo,x.tema,x.descripcion,x.estado,x.prioridad,x.proxima_fecha.isoformat() if x.proxima_fecha else None,x.responsable,sid)); cur.commit(); return {'ok':True}

@app.get('/api/reuniones-familia')
def get_reuniones(student_id:Optional[int]=None,u=Depends(current_user)):
    if student_id is not None and not can_access_student(u,student_id): raise HTTPException(403,'Sin acceso al estudiante')
    q='SELECT r.*,e.matricula,e.nombre estudiante FROM reuniones_familia r JOIN estudiantes e ON e.id=r.estudiante_id WHERE 1=1'; args=[]
    if u['rol']!='admin':
        ids=allowed_course_ids(u) or []
        if not ids:return []
        q+=' AND e.curso_id IN ('+','.join('?'*len(ids))+')'; args.extend(ids)
    if student_id is not None:q+=' AND r.estudiante_id=?';args.append(student_id)
    q+=' ORDER BY r.fecha DESC,r.id DESC';return rows(q,args)

@app.post('/api/reuniones-familia')
def post_reunion(x:ReunionIn,u=Depends(current_user)):
    if not can_access_student(u,x.estudiante_id): raise HTTPException(403,'Sin acceso al estudiante')
    cur=conn();cur.execute('INSERT INTO reuniones_familia(estudiante_id,fecha,hora,tipo,participantes,acuerdos,observaciones,fecha_registro) VALUES(?,?,?,?,?,?,?,?)',(x.estudiante_id,x.fecha.isoformat(),x.hora,x.tipo,x.participantes,x.acuerdos,x.observaciones,now()));cur.commit();return {'id':cur.lastrowid}

@app.get('/api/compromisos')
def get_compromisos(student_id:Optional[int]=None,estado:Optional[str]=None,u=Depends(current_user)):
    if student_id is not None and not can_access_student(u,student_id): raise HTTPException(403,'Sin acceso al estudiante')
    q='SELECT c.*,e.matricula,e.nombre estudiante FROM compromisos c JOIN estudiantes e ON e.id=c.estudiante_id WHERE 1=1';args=[]
    if u['rol']!='admin':
        ids=allowed_course_ids(u) or []
        if not ids:return []
        q+=' AND e.curso_id IN ('+','.join('?'*len(ids))+')';args.extend(ids)
    if student_id is not None:q+=' AND c.estudiante_id=?';args.append(student_id)
    if estado:q+=' AND c.estado=?';args.append(estado)
    q+=' ORDER BY CASE WHEN c.estado="Pendiente" THEN 0 ELSE 1 END,c.fecha_limite ASC,c.id DESC';return rows(q,args)

@app.post('/api/compromisos')
def post_compromiso(x:CompromisoIn,u=Depends(current_user)):
    if not can_access_student(u,x.estudiante_id): raise HTTPException(403,'Sin acceso al estudiante')
    if x.seguimiento_id:
        r=rows('SELECT estudiante_id FROM seguimientos WHERE id=?',(x.seguimiento_id,))
        if not r or r[0]['estudiante_id']!=x.estudiante_id: raise HTTPException(400,'Seguimiento no válido')
    cur=conn();cur.execute('INSERT INTO compromisos(estudiante_id,seguimiento_id,descripcion,fecha_limite,estado,responsable,fecha_cumplimiento,observacion,fecha_registro) VALUES(?,?,?,?,?,?,?,?,?)',(x.estudiante_id,x.seguimiento_id,x.descripcion,x.fecha_limite.isoformat() if x.fecha_limite else None,x.estado,x.responsable,x.fecha_cumplimiento.isoformat() if x.fecha_cumplimiento else None,x.observacion,now()));cur.commit();return {'id':cur.lastrowid}

@app.put('/api/compromisos/{cid}')
def put_compromiso(cid:int,x:CompromisoIn,u=Depends(current_user)):
    r=rows('SELECT estudiante_id FROM compromisos WHERE id=?',(cid,));
    if not r:raise HTTPException(404,'Compromiso no encontrado')
    if not can_access_student(u,r[0]['estudiante_id']):raise HTTPException(403,'Sin acceso')
    cur=conn();cur.execute('UPDATE compromisos SET descripcion=?,fecha_limite=?,estado=?,responsable=?,fecha_cumplimiento=?,observacion=? WHERE id=?',(x.descripcion,x.fecha_limite.isoformat() if x.fecha_limite else None,x.estado,x.responsable,x.fecha_cumplimiento.isoformat() if x.fecha_cumplimiento else None,x.observacion,cid));cur.commit();return {'ok':True}

@app.get('/api/usuarios')
def usuarios(u=Depends(admin_user)):
    return rows("SELECT u.id,u.usuario,u.nombre,u.rol,u.activo,u.fecha_registro,u.docente_id,d.nombre docente_nombre,d.especialidad docente_especialidad FROM usuarios u LEFT JOIN docentes d ON d.id=u.docente_id ORDER BY u.nombre COLLATE NOCASE")

@app.post('/api/usuarios')
def crear_usuario(x:UsuarioCreate,u=Depends(admin_user)):
    if x.rol not in ('admin','docente'): raise HTTPException(400,'Rol inválido')
    c=conn()
    try:
        cur=c.execute('INSERT INTO usuarios(usuario,nombre,rol,password_hash,activo,fecha_registro) VALUES(?,?,?,?,1,?)',(x.usuario.strip(),x.nombre.strip(),x.rol,hash_password(x.password),now())); c.commit(); r=dict(c.execute('SELECT id,usuario,nombre,rol,activo,fecha_registro FROM usuarios WHERE id=?',(cur.lastrowid,)).fetchone())
    except sqlite3.IntegrityError: c.close(); raise HTTPException(409,'El usuario ya existe')
    c.close(); return r

@app.put('/api/usuarios/{usuario_id}')
def actualizar_usuario(usuario_id:int,x:UsuarioUpdate,u=Depends(admin_user)):
    if x.rol not in ('admin','docente'): raise HTTPException(400,'Rol inválido')
    c=conn(); r=c.execute('SELECT * FROM usuarios WHERE id=?',(usuario_id,)).fetchone()
    if not r: c.close(); raise HTTPException(404,'Usuario no encontrado')
    if usuario_id==u['id'] and not x.activo: c.close(); raise HTTPException(400,'No puedes desactivar tu propia cuenta')
    if r['rol']=='admin' and x.rol!='admin' and c.execute("SELECT COUNT(*) FROM usuarios WHERE rol='admin' AND activo=1").fetchone()[0] <= 1: c.close(); raise HTTPException(400,'Debe existir al menos un administrador activo')
    c.execute('UPDATE usuarios SET nombre=?,rol=?,activo=? WHERE id=?',(x.nombre.strip(),x.rol,int(x.activo),usuario_id)); c.commit(); out=dict(c.execute('SELECT id,usuario,nombre,rol,activo,fecha_registro FROM usuarios WHERE id=?',(usuario_id,)).fetchone()); c.close(); return out

@app.put('/api/usuarios/{usuario_id}/docente')
def vincular_usuario_docente(usuario_id:int,x:VincularDocenteUsuarioIn,u=Depends(admin_user)):
    c=conn(); r=c.execute('SELECT * FROM usuarios WHERE id=?',(usuario_id,)).fetchone()
    if not r: c.close(); raise HTTPException(404,'Usuario no encontrado')
    if r['rol']!='docente' and x.docente_id is not None: c.close(); raise HTTPException(400,'Solo los usuarios con rol docente pueden vincularse a un docente')
    if x.docente_id is not None:
        d=c.execute("SELECT id FROM docentes WHERE id=? AND estado='Activo'",(x.docente_id,)).fetchone()
        if not d: c.close(); raise HTTPException(400,'El docente seleccionado no está activo o no existe')
    c.execute('UPDATE usuarios SET docente_id=? WHERE id=?',(x.docente_id,usuario_id)); c.commit(); out=dict(c.execute('SELECT u.id,u.usuario,u.nombre,u.rol,u.activo,u.fecha_registro,u.docente_id,d.nombre docente_nombre,d.especialidad docente_especialidad FROM usuarios u LEFT JOIN docentes d ON d.id=u.docente_id WHERE u.id=?',(usuario_id,)).fetchone()); c.close(); return out

@app.put('/api/usuarios/{usuario_id}/password')
def cambiar_password(usuario_id:int,x:PasswordUpdate,u=Depends(current_user)):
    if usuario_id!=u['id'] and u['rol']!='admin': raise HTTPException(403,'Solo puedes cambiar tu propia contraseña')
    c=conn(); r=c.execute('SELECT id FROM usuarios WHERE id=?',(usuario_id,)).fetchone()
    if not r: c.close(); raise HTTPException(404,'Usuario no encontrado')
    c.execute('UPDATE usuarios SET password_hash=? WHERE id=?',(hash_password(x.password),usuario_id)); c.commit(); c.close(); return {'ok':True}


@app.get('/api/competencias')
def listar_competencias(asignatura_id:Optional[int]=None,u=Depends(current_user)):
    sql="SELECT c.*,a.nombre asignatura_nombre FROM competencias c JOIN asignaturas a ON a.id=c.asignatura_id WHERE c.activa=1"; args=[]
    if asignatura_id: sql += " AND c.asignatura_id=?"; args.append(asignatura_id)
    if u['rol']!='admin':
        ids=sorted(set(p['asignatura_id'] for p in (allowed_subject_pairs(u) or [])))
        if not ids:return []
        sql += ' AND c.asignatura_id IN ('+','.join('?'*len(ids))+')'; args.extend(ids)
    sql += " ORDER BY c.asignatura_id,c.orden,c.codigo"; return rows(sql,args)
@app.post('/api/competencias')
def crear_competencia(x:CompetenciaIn,u=Depends(current_user)):
    if not rows("SELECT id FROM asignaturas WHERE id=? AND activa=1",(x.asignatura_id,)): raise HTTPException(404,'Asignatura no encontrada o inactiva')
    c=conn()
    try:
        cur=c.execute('INSERT INTO competencias(asignatura_id,codigo,nombre,descripcion,orden,activa) VALUES(?,?,?,?,?,?)',(x.asignatura_id,x.codigo.strip(),x.nombre.strip(),x.descripcion or '',x.orden,int(x.activa)));c.commit(); out=dict(c.execute('SELECT * FROM competencias WHERE id=?',(cur.lastrowid,)).fetchone())
    except sqlite3.IntegrityError: c.close(); raise HTTPException(409,'El código de competencia ya existe para esa asignatura')
    c.close(); return out
@app.put('/api/competencias/{competencia_id}')
def actualizar_competencia(competencia_id:int,x:CompetenciaIn,u=Depends(current_user)):
    c=conn(); r=c.execute('SELECT id FROM competencias WHERE id=?',(competencia_id,)).fetchone()
    if not r: c.close(); raise HTTPException(404,'Competencia no encontrada')
    try:
        c.execute('UPDATE competencias SET asignatura_id=?,codigo=?,nombre=?,descripcion=?,orden=?,activa=? WHERE id=?',(x.asignatura_id,x.codigo.strip(),x.nombre.strip(),x.descripcion or '',x.orden,int(x.activa),competencia_id));c.commit();out=dict(c.execute('SELECT * FROM competencias WHERE id=?',(competencia_id,)).fetchone())
    except sqlite3.IntegrityError: c.close(); raise HTTPException(409,'El código de competencia ya existe para esa asignatura')
    c.close(); return out
@app.get('/api/competencias-fundamentales')
def listar_fundamentales(u=Depends(current_user)):
    return rows('SELECT * FROM competencias_fundamentales WHERE activa=1 ORDER BY orden,codigo')
@app.post('/api/competencias-fundamentales')
def crear_fundamental(x:FundamentalIn,u=Depends(current_user)):
    c=conn()
    try: cur=c.execute('INSERT INTO competencias_fundamentales(codigo,nombre,descripcion,orden,activa) VALUES(?,?,?,?,?)',(x.codigo.strip(),x.nombre.strip(),x.descripcion or '',x.orden,int(x.activa)));c.commit();out=dict(c.execute('SELECT * FROM competencias_fundamentales WHERE id=?',(cur.lastrowid,)).fetchone())
    except sqlite3.IntegrityError: c.close();raise HTTPException(409,'El código de competencia fundamental ya existe')
    c.close();return out
@app.post('/api/competencias-fundamentales/mapear')
def mapear_competencia(x:MapeoIn,u=Depends(current_user)):
    if not rows('SELECT id FROM competencias_fundamentales WHERE id=?',(x.fundamental_id,)): raise HTTPException(404,'Competencia fundamental no encontrada')
    if not rows('SELECT id FROM competencias WHERE id=?',(x.especifica_id,)): raise HTTPException(404,'Competencia específica no encontrada')
    c=conn();c.execute('INSERT OR IGNORE INTO competencia_fundamental_especifica(fundamental_id,especifica_id) VALUES(?,?)',(x.fundamental_id,x.especifica_id));c.commit();c.close();return {'ok':True}
@app.get('/api/competencias-fundamentales/{fundamental_id}/especificas')
def fundamentales_especificas(fundamental_id:int,u=Depends(current_user)):
    return rows('SELECT c.* FROM competencias c JOIN competencia_fundamental_especifica m ON m.especifica_id=c.id WHERE m.fundamental_id=? ORDER BY c.orden,c.codigo',(fundamental_id,))
@app.get('/api/actividades')
def listar_actividades(estudiante_id:Optional[int]=None,asignatura_id:Optional[int]=None,periodo:Optional[int]=None,u=Depends(current_user)):
    sql='''SELECT ac.*,e.nombre estudiante_nombre,a.nombre asignatura_nombre,c.codigo competencia_codigo,c.nombre competencia_nombre FROM actividades ac JOIN estudiantes e ON e.id=ac.estudiante_id JOIN asignaturas a ON a.id=ac.asignatura_id JOIN competencias c ON c.id=ac.competencia_id WHERE 1=1''';args=[]
    if estudiante_id:sql+=' AND ac.estudiante_id=?';args.append(estudiante_id)
    if asignatura_id:sql+=' AND ac.asignatura_id=?';args.append(asignatura_id)
    if periodo:sql+=' AND ac.periodo=?';args.append(periodo)
    
    if u['rol']!='admin':
        did=u.get('docente_id'); sql+=' AND EXISTS (SELECT 1 FROM cursos cx LEFT JOIN curso_asignatura cax ON cax.curso_id=cx.id AND cax.activa=1 WHERE cx.id=(SELECT curso_id FROM estudiantes WHERE id=ac.estudiante_id) AND (cx.docente_id=? OR cax.docente_id=?))'; args.extend([did,did])
    sql+=' ORDER BY ac.fecha DESC,ac.id DESC';return rows(sql,args)
@app.post('/api/actividades')
def crear_actividad(x:ActividadIn,u=Depends(current_user)):
    if not rows('SELECT id FROM estudiantes WHERE id=?',(x.estudiante_id,)):raise HTTPException(404,'Estudiante no encontrado')
    if not can_access_student(u,x.estudiante_id): raise HTTPException(403,'No tienes acceso a este estudiante')
    if not rows('SELECT id FROM competencias WHERE id=? AND asignatura_id=?',(x.competencia_id,x.asignatura_id)):raise HTTPException(400,'La competencia no pertenece a la asignatura seleccionada')
    c=conn();cur=c.execute('INSERT INTO actividades(estudiante_id,asignatura_id,competencia_id,periodo,fecha,titulo,descripcion,tipo_evaluacion,instrumento,calificacion,observacion,fecha_registro) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',(x.estudiante_id,x.asignatura_id,x.competencia_id,x.periodo,x.fecha.isoformat(),x.titulo.strip(),x.descripcion,x.tipo_evaluacion,x.instrumento,x.calificacion,x.observacion,now()));c.commit();out=dict(c.execute('SELECT * FROM actividades WHERE id=?',(cur.lastrowid,)).fetchone());c.close();return out
@app.get('/api/libro-calificaciones')
def libro_calificaciones(curso_id:int,asignatura_id:int,periodo:int=1,u=Depends(current_user)):
    if not can_access_course(u,curso_id): raise HTTPException(403,'No tienes acceso a esta clase')
    if not rows('SELECT id FROM curso_asignatura WHERE curso_id=? AND asignatura_id=? AND activa=1',(curso_id,asignatura_id)): raise HTTPException(400,'La asignatura no está asignada a la clase')
    students=rows("SELECT id,nombre,matricula,orden_lista FROM estudiantes WHERE curso_id=? AND estado='Activo' ORDER BY nombre,matricula",(curso_id,))
    comps=rows('SELECT id,codigo,nombre,orden FROM competencias WHERE asignatura_id=? AND activa=1 ORDER BY orden,codigo',(asignatura_id,))
    vals=rows('SELECT estudiante_id,competencia_id,tipo,calificacion,observacion FROM calificaciones WHERE asignatura_id=? AND periodo=? AND estudiante_id IN (SELECT id FROM estudiantes WHERE curso_id=?)',(asignatura_id,periodo,curso_id))
    by={(r['estudiante_id'],r['competencia_id'],r['tipo']):r for r in vals}
    out=[]
    for st in students:
        row={'estudiante_id':st['id'],'estudiante':st['nombre'],'matricula':st['matricula'],'competencias':[]}
        for c in comps:
            r=by.get((st['id'],c['id'],'P'))
            row['competencias'].append({'competencia_id':c['id'],'codigo':c['codigo'],'nombre':c['nombre'],'calificacion':r['calificacion'] if r else None,'observacion':r['observacion'] if r else None})
        out.append(row)
    return {'curso_id':curso_id,'asignatura_id':asignatura_id,'periodo':periodo,'competencias':comps,'filas':out}

@app.post('/api/calificaciones/lote')
def guardar_calificaciones_lote(x:CalificacionLoteIn,u=Depends(current_user)):
    if not can_access_course(u,x.curso_id): raise HTTPException(403,'No tienes acceso a esta clase')
    pair=rows('SELECT docente_id FROM curso_asignatura WHERE curso_id=? AND asignatura_id=? AND activa=1',(x.curso_id,x.asignatura_id))
    if not pair: raise HTTPException(400,'La asignatura no está asignada a la clase')
    if u['rol']!='admin' and pair[0]['docente_id']!=u.get('docente_id') and not rows('SELECT id FROM cursos WHERE id=? AND docente_id=?',(x.curso_id,u.get('docente_id'))): raise HTTPException(403,'No puedes calificar esta asignatura')
    c=conn(); n=0
    for r in x.registros:
        if not rows('SELECT id FROM estudiantes WHERE id=? AND curso_id=?',(r.estudiante_id,x.curso_id)): raise HTTPException(400,f'El estudiante {r.estudiante_id} no pertenece a la clase')
        if not rows('SELECT id FROM competencias WHERE id=? AND asignatura_id=?',(r.competencia_id,x.asignatura_id)): raise HTTPException(400,'Competencia inválida para la asignatura')
        c.execute('''INSERT INTO calificaciones(estudiante_id,asignatura_id,competencia_id,periodo,tipo,calificacion,observacion,fecha_actualizacion) VALUES(?,?,?,?,?,?,?,?) ON CONFLICT(estudiante_id,asignatura_id,competencia_id,periodo,tipo) DO UPDATE SET calificacion=excluded.calificacion,observacion=excluded.observacion,fecha_actualizacion=excluded.fecha_actualizacion''',(r.estudiante_id,x.asignatura_id,r.competencia_id,x.periodo,x.tipo,r.calificacion,r.observacion,now())); n+=1
    c.commit();c.close();return {'ok':True,'guardadas':n}

@app.put('/api/actividades/{actividad_id}')
def actualizar_actividad(actividad_id:int,x:ActividadUpdateIn,u=Depends(current_user)):
    old=rows('SELECT * FROM actividades WHERE id=?',(actividad_id,))
    if not old: raise HTTPException(404,'Actividad no encontrada')
    if not can_access_student(u,old[0]['estudiante_id']): raise HTTPException(403,'No tienes acceso a esta actividad')
    c=conn();c.execute('UPDATE actividades SET fecha=?,titulo=?,descripcion=?,tipo_evaluacion=?,instrumento=?,calificacion=?,observacion=? WHERE id=?',(x.fecha.isoformat(),x.titulo.strip(),x.descripcion,x.tipo_evaluacion,x.instrumento,x.calificacion,x.observacion,actividad_id));c.commit();out=dict(c.execute('SELECT * FROM actividades WHERE id=?',(actividad_id,)).fetchone());c.close();return out

@app.delete('/api/actividades/{actividad_id}')
def eliminar_actividad(actividad_id:int,u=Depends(current_user)):
    old=rows('SELECT * FROM actividades WHERE id=?',(actividad_id,))
    if not old: raise HTTPException(404,'Actividad no encontrada')
    if not can_access_student(u,old[0]['estudiante_id']): raise HTTPException(403,'No tienes acceso a esta actividad')
    c=conn();c.execute('DELETE FROM actividades WHERE id=?',(actividad_id,));c.commit();c.close();return {'ok':True}

@app.get('/api/calificaciones')
def listar_calificaciones(estudiante_id:Optional[int]=None,asignatura_id:Optional[int]=None,periodo:Optional[int]=None,u=Depends(current_user)):
    sql='''SELECT c.*,e.nombre estudiante_nombre,a.nombre asignatura_nombre,co.codigo competencia_codigo,co.nombre competencia_nombre FROM calificaciones c JOIN estudiantes e ON e.id=c.estudiante_id JOIN asignaturas a ON a.id=c.asignatura_id JOIN competencias co ON co.id=c.competencia_id WHERE 1=1''';args=[]
    if estudiante_id:sql+=' AND c.estudiante_id=?';args.append(estudiante_id)
    if asignatura_id:sql+=' AND c.asignatura_id=?';args.append(asignatura_id)
    if periodo:sql+=' AND c.periodo=?';args.append(periodo)
    
    if u['rol']!='admin':
        did=u.get('docente_id'); sql+=' AND EXISTS (SELECT 1 FROM cursos cx LEFT JOIN curso_asignatura cax ON cax.curso_id=cx.id AND cax.activa=1 WHERE cx.id=(SELECT curso_id FROM estudiantes WHERE id=c.estudiante_id) AND (cx.docente_id=? OR cax.docente_id=?))'; args.extend([did,did])
    sql+=' ORDER BY e.nombre, a.nombre, co.orden, c.periodo, c.tipo';return rows(sql,args)
@app.post('/api/calificaciones')
def guardar_calificacion(x:CalificacionIn,u=Depends(current_user)):
    if not rows('SELECT id FROM estudiantes WHERE id=?',(x.estudiante_id,)):raise HTTPException(404,'Estudiante no encontrado')
    if not can_access_student(u,x.estudiante_id): raise HTTPException(403,'No tienes acceso a este estudiante')
    if not rows('SELECT id FROM competencias WHERE id=? AND asignatura_id=?',(x.competencia_id,x.asignatura_id)):raise HTTPException(400,'La competencia no pertenece a la asignatura seleccionada')
    c=conn();c.execute('''INSERT INTO calificaciones(estudiante_id,asignatura_id,competencia_id,periodo,tipo,calificacion,observacion,fecha_actualizacion) VALUES(?,?,?,?,?,?,?,?) ON CONFLICT(estudiante_id,asignatura_id,competencia_id,periodo,tipo) DO UPDATE SET calificacion=excluded.calificacion,observacion=excluded.observacion,fecha_actualizacion=excluded.fecha_actualizacion''',(x.estudiante_id,x.asignatura_id,x.competencia_id,x.periodo,x.tipo,x.calificacion,x.observacion,now()));c.commit();out=dict(c.execute('SELECT * FROM calificaciones WHERE estudiante_id=? AND asignatura_id=? AND competencia_id=? AND periodo=? AND tipo=?',(x.estudiante_id,x.asignatura_id,x.competencia_id,x.periodo,x.tipo)).fetchone());c.close();return out
@app.get('/api/rendimiento/estudiante/{estudiante_id}')
def rendimiento(estudiante_id:int,u=Depends(current_user)):
    if not can_access_student(u,estudiante_id): raise HTTPException(403,'No tienes acceso a este estudiante')
    acts=rows('''SELECT a.id,a.nombre asignatura, c.id competencia_id,c.codigo competencia_codigo,c.nombre competencia, ac.periodo, ROUND(AVG(ac.calificacion),2) promedio, COUNT(ac.id) actividades FROM actividades ac JOIN asignaturas a ON a.id=ac.asignatura_id JOIN competencias c ON c.id=ac.competencia_id WHERE ac.estudiante_id=? GROUP BY a.id,c.id,ac.periodo ORDER BY a.nombre,c.orden,ac.periodo''',(estudiante_id,))
    cal=rows('''SELECT a.nombre asignatura,c.codigo competencia_codigo,c.nombre competencia,cl.periodo,cl.tipo,cl.calificacion FROM calificaciones cl JOIN asignaturas a ON a.id=cl.asignatura_id JOIN competencias c ON c.id=cl.competencia_id WHERE cl.estudiante_id=? ORDER BY a.nombre,c.orden,cl.periodo,cl.tipo''',(estudiante_id,))
    return {'actividades':acts,'calificaciones':cal}

@app.get('/api/mi-carga')
def mi_carga(u=Depends(current_user)):
    if u['rol']=='admin':
        return {'rol':'admin','docente_id':None,'clases':rows('SELECT c.id,c.nombre,c.grado,c.seccion,c.anio_escolar,d.nombre docente_tutor FROM cursos c LEFT JOIN docentes d ON d.id=c.docente_id ORDER BY c.nombre'), 'asignaciones':rows('SELECT ca.curso_id,ca.asignatura_id,c.nombre curso_nombre,a.nombre asignatura,d.nombre docente_nombre,d.especialidad FROM curso_asignatura ca JOIN cursos c ON c.id=ca.curso_id JOIN asignaturas a ON a.id=ca.asignatura_id LEFT JOIN docentes d ON d.id=ca.docente_id WHERE ca.activa=1 ORDER BY c.nombre,a.nombre')}
    did=u.get('docente_id')
    if not did:return {'rol':'docente','docente_id':None,'clases':[],'asignaciones':[]}
    return {'rol':'docente','docente_id':did,'clases':rows('SELECT c.id,c.nombre,c.grado,c.seccion,c.anio_escolar,d.nombre docente_tutor FROM cursos c LEFT JOIN docentes d ON d.id=c.docente_id WHERE c.docente_id=? OR EXISTS (SELECT 1 FROM curso_asignatura ca WHERE ca.curso_id=c.id AND ca.docente_id=? AND ca.activa=1) ORDER BY c.nombre',(did,did)), 'asignaciones':rows('SELECT ca.curso_id,ca.asignatura_id,c.nombre curso_nombre,a.nombre asignatura,d.nombre docente_nombre,d.especialidad FROM curso_asignatura ca JOIN cursos c ON c.id=ca.curso_id JOIN asignaturas a ON a.id=ca.asignatura_id LEFT JOIN docentes d ON d.id=ca.docente_id WHERE ca.docente_id=? AND ca.activa=1 ORDER BY c.nombre,a.nombre',(did,))}


def can_access_course(u,curso_id):
    return u['rol']=='admin' or curso_id in (allowed_course_ids(u) or [])

def student_course(student_id):
    r=rows('SELECT curso_id FROM estudiantes WHERE id=?',(student_id,))
    return r[0]['curso_id'] if r else None

@app.get('/api/asistencia')
def listar_asistencia(fecha:Optional[date]=None,curso_id:Optional[int]=None,estudiante_id:Optional[int]=None,u=Depends(current_user)):
    sql="""SELECT a.*,e.nombre estudiante_nombre,e.matricula,c.id curso_id,c.nombre curso_nombre,c.grado,c.seccion
           FROM asistencia a JOIN estudiantes e ON e.id=a.estudiante_id LEFT JOIN cursos c ON c.id=e.curso_id WHERE 1=1"""; args=[]
    if fecha: sql+=' AND a.fecha=?'; args.append(fecha.isoformat())
    if curso_id: sql+=' AND e.curso_id=?'; args.append(curso_id)
    if estudiante_id: sql+=' AND e.id=?'; args.append(estudiante_id)
    if u['rol']!='admin':
        ids=allowed_course_ids(u) or []
        if not ids:return []
        sql+=' AND e.curso_id IN ('+','.join('?'*len(ids))+')'; args.extend(ids)
    sql+=' ORDER BY a.fecha DESC,e.nombre COLLATE NOCASE'; return rows(sql,args)

@app.post('/api/asistencia')
def guardar_asistencia(x:AsistenciaIn,u=Depends(current_user)):
    cid=student_course(x.estudiante_id)
    if cid is None: raise HTTPException(404,'Estudiante no encontrado')
    if not can_access_course(u,cid): raise HTTPException(403,'No tienes acceso a este estudiante')
    if x.estado not in {'Presente','Ausente','Tardanza','Excusa'}: raise HTTPException(400,'Estado de asistencia inválido')
    c=conn(); c.execute("""INSERT INTO asistencia(estudiante_id,fecha,estado,hora_entrada,observacion) VALUES(?,?,?,?,?)
      ON CONFLICT(estudiante_id,fecha) DO UPDATE SET estado=excluded.estado,hora_entrada=excluded.hora_entrada,observacion=excluded.observacion""",(x.estudiante_id,x.fecha.isoformat(),x.estado,x.hora_entrada,x.observacion)); c.commit()
    out=dict(c.execute('SELECT * FROM asistencia WHERE estudiante_id=? AND fecha=?',(x.estudiante_id,x.fecha.isoformat())).fetchone()); c.close(); return out

@app.post('/api/asistencia/lote')
def guardar_asistencia_lote(x:AsistenciaLoteIn,u=Depends(current_user)):
    if not can_access_course(u,x.curso_id): raise HTTPException(403,'No tienes acceso a esta clase')
    c=conn(); count=0
    for item in x.registros:
        sid=int(item.get('estudiante_id')); estado=str(item.get('estado','Presente')); hora=item.get('hora_entrada'); obs=item.get('observacion')
        r=c.execute('SELECT id FROM estudiantes WHERE id=? AND curso_id=?',(sid,x.curso_id)).fetchone()
        if not r or estado not in {'Presente','Ausente','Tardanza','Excusa'}: continue
        c.execute("""INSERT INTO asistencia(estudiante_id,fecha,estado,hora_entrada,observacion) VALUES(?,?,?,?,?)
          ON CONFLICT(estudiante_id,fecha) DO UPDATE SET estado=excluded.estado,hora_entrada=excluded.hora_entrada,observacion=excluded.observacion""",(sid,x.fecha.isoformat(),estado,hora,obs)); count+=1
    c.commit(); c.close(); return {'ok':True,'guardados':count,'fecha':x.fecha.isoformat(),'curso_id':x.curso_id}

@app.get('/api/asistencia/resumen')
def resumen_asistencia(curso_id:int,fecha_desde:Optional[date]=None,fecha_hasta:Optional[date]=None,u=Depends(current_user)):
    if not can_access_course(u,curso_id): raise HTTPException(403,'No tienes acceso a esta clase')
    join='a.estudiante_id=e.id AND a.fecha>=COALESCE(?,a.fecha) AND a.fecha<=COALESCE(?,a.fecha)'; args=[fecha_desde.isoformat() if fecha_desde else None,fecha_hasta.isoformat() if fecha_hasta else None,curso_id]
    return rows(f"""SELECT e.id,e.matricula,e.nombre,COUNT(a.id) total_registros,
      SUM(CASE WHEN a.estado='Presente' THEN 1 ELSE 0 END) presentes,
      SUM(CASE WHEN a.estado='Ausente' THEN 1 ELSE 0 END) ausencias,
      SUM(CASE WHEN a.estado='Tardanza' THEN 1 ELSE 0 END) tardanzas,
      SUM(CASE WHEN a.estado='Excusa' THEN 1 ELSE 0 END) excusas
      FROM estudiantes e LEFT JOIN asistencia a ON {join} WHERE e.curso_id=? GROUP BY e.id ORDER BY e.nombre COLLATE NOCASE""",args)

@app.get('/api/incidencias')
def listar_incidencias(estudiante_id:Optional[int]=None,curso_id:Optional[int]=None,u=Depends(current_user)):
    sql="""SELECT i.*,e.nombre estudiante_nombre,e.matricula,e.curso_id,c.nombre curso_nombre FROM incidencias i JOIN estudiantes e ON e.id=i.estudiante_id LEFT JOIN cursos c ON c.id=e.curso_id WHERE 1=1"""; args=[]
    if estudiante_id: sql+=' AND i.estudiante_id=?'; args.append(estudiante_id)
    if curso_id: sql+=' AND e.curso_id=?'; args.append(curso_id)
    if u['rol']!='admin':
        ids=allowed_course_ids(u) or []
        if not ids:return []
        sql+=' AND e.curso_id IN ('+','.join('?'*len(ids))+')'; args.extend(ids)
    sql+=' ORDER BY i.fecha DESC,i.id DESC'; return rows(sql,args)

@app.post('/api/incidencias')
def crear_incidencia(x:IncidenciaIn,u=Depends(current_user)):
    cid=student_course(x.estudiante_id)
    if cid is None: raise HTTPException(404,'Estudiante no encontrado')
    if not can_access_course(u,cid): raise HTTPException(403,'No tienes acceso a este estudiante')
    c=conn(); cur=c.execute('INSERT INTO incidencias(estudiante_id,fecha,hora,tipo,gravedad,descripcion,accion_tomada,observacion) VALUES(?,?,?,?,?,?,?,?)',(x.estudiante_id,x.fecha.isoformat(),x.hora,x.tipo,x.gravedad,x.descripcion,x.accion_tomada,x.observacion)); c.commit(); out=dict(c.execute('SELECT * FROM incidencias WHERE id=?',(cur.lastrowid,)).fetchone()); c.close(); return out

@app.get('/api/permisos')
def listar_permisos(estudiante_id:Optional[int]=None,curso_id:Optional[int]=None,estado:Optional[str]=None,u=Depends(current_user)):
    sql="""SELECT p.*,e.nombre estudiante_nombre,e.matricula,e.curso_id,c.nombre curso_nombre FROM permisos p JOIN estudiantes e ON e.id=p.estudiante_id LEFT JOIN cursos c ON c.id=e.curso_id WHERE 1=1"""; args=[]
    if estudiante_id: sql+=' AND p.estudiante_id=?'; args.append(estudiante_id)
    if curso_id: sql+=' AND e.curso_id=?'; args.append(curso_id)
    if estado: sql+=' AND p.estado=?'; args.append(estado)
    if u['rol']!='admin':
        ids=allowed_course_ids(u) or []
        if not ids:return []
        sql+=' AND e.curso_id IN ('+','.join('?'*len(ids))+')'; args.extend(ids)
    sql+=' ORDER BY p.fecha DESC,p.id DESC'; return rows(sql,args)

@app.post('/api/permisos')
def crear_permiso(x:PermisoIn,u=Depends(current_user)):
    cid=student_course(x.estudiante_id)
    if cid is None: raise HTTPException(404,'Estudiante no encontrado')
    if not can_access_course(u,cid): raise HTTPException(403,'No tienes acceso a este estudiante')
    c=conn(); cur=c.execute('INSERT INTO permisos(estudiante_id,fecha,tipo,hora_desde,hora_hasta,motivo,estado,observacion) VALUES(?,?,?,?,?,?,?,?)',(x.estudiante_id,x.fecha.isoformat(),x.tipo,x.hora_desde,x.hora_hasta,x.motivo,x.estado,x.observacion)); c.commit(); out=dict(c.execute('SELECT * FROM permisos WHERE id=?',(cur.lastrowid,)).fetchone()); c.close(); return out

@app.put('/api/permisos/{permiso_id}')
def actualizar_permiso(permiso_id:int,x:PermisoUpdateIn,u=Depends(current_user)):
    r=rows('SELECT estudiante_id FROM permisos WHERE id=?',(permiso_id,))
    if not r: raise HTTPException(404,'Permiso no encontrado')
    cid=student_course(r[0]['estudiante_id'])
    if not can_access_course(u,cid): raise HTTPException(403,'No tienes acceso a este permiso')
    if x.estado is not None and x.estado not in {'Pendiente','Aprobado','Rechazado'}: raise HTTPException(400,'Estado de permiso inválido')
    current=rows('SELECT * FROM permisos WHERE id=?',(permiso_id,))[0]
    fecha=(x.fecha.isoformat() if x.fecha is not None else current['fecha'])
    tipo=x.tipo if x.tipo is not None else current['tipo']
    estado=x.estado if x.estado is not None else current['estado']
    c=conn(); c.execute('UPDATE permisos SET fecha=?,tipo=?,hora_desde=?,hora_hasta=?,motivo=?,estado=?,observacion=? WHERE id=?',(fecha,tipo,x.hora_desde if x.hora_desde is not None else current['hora_desde'],x.hora_hasta if x.hora_hasta is not None else current['hora_hasta'],x.motivo if x.motivo is not None else current['motivo'],estado,x.observacion if x.observacion is not None else current['observacion'],permiso_id)); c.commit(); out=dict(c.execute('SELECT * FROM permisos WHERE id=?',(permiso_id,)).fetchone()); c.close(); return out

# ========================= ETAPA 6: CIERRE ACADÉMICO =========================
def _cfg(key, default=None):
    r=rows('SELECT valor FROM configuracion WHERE clave=?',(key,))
    return r[0]['valor'] if r else default

def _periodo_competencia(sid,aid,cid,p):
    r=rows('SELECT ROUND(AVG(calificacion),2) v FROM actividades WHERE estudiante_id=? AND asignatura_id=? AND competencia_id=? AND periodo=?',(sid,aid,cid,p))
    return r[0]['v'] if r and r[0]['v'] is not None else None

def _rp_competencia(sid,aid,cid,p):
    r=rows("SELECT calificacion FROM calificaciones WHERE estudiante_id=? AND asignatura_id=? AND competencia_id=? AND periodo=? AND tipo='RP'",(sid,aid,cid,p))
    return r[0]['calificacion'] if r and r[0]['calificacion'] is not None else None

def _competencia_anual(sid,aid,cid):
    vals=[]; use_rp=_cfg('rp_reemplaza_p','1')=='1'
    for p in range(1,5):
        pv=_periodo_competencia(sid,aid,cid,p); rv=_rp_competencia(sid,aid,cid,p)
        if use_rp and rv is not None: vals.append(rv)
        elif pv is not None: vals.append(pv)
        elif rv is not None: vals.append(rv)
    return round(sum(vals)/len(vals),2) if vals else None

def _calificacion_final_base(sid,aid):
    cs=rows('SELECT id FROM competencias WHERE asignatura_id=? AND activa=1 ORDER BY orden,codigo',(aid,))
    vals=[_competencia_anual(sid,aid,int(c['id'])) for c in cs]
    vals=[v for v in vals if v is not None]
    return round(sum(vals)/len(vals),2) if vals else None

def _resultado_final(cf, r):
    wcf=float(_cfg('peso_completiva_cf','50'))/100; wcec=float(_cfg('peso_completiva_cec','50'))/100
    wecf=float(_cfg('peso_extra_cf','30'))/100; wecex=float(_cfg('peso_extra_ceex','70'))/100
    cec=r.get('completiva'); ceex=r.get('extraordinaria'); ecf=r.get('especial_cf'); ece=r.get('especial_ce')
    ccf=round(cf*wcf+cec*wcec,2) if cf is not None and cec is not None else None
    cexf=round(cf*wecf+ceex*wecex,2) if cf is not None and ceex is not None else None
    final=ece if ece is not None else (ecf if ecf is not None else (cexf if cexf is not None else (ccf if ccf is not None else cf)))
    umbral=float(_cfg('umbral_aprobacion','70'))
    sit=(r.get('situacion') or ('A' if final is not None and final>=umbral else ('R' if final is not None else ''))).upper()
    return {'CF':cf,'CCF':ccf,'CEXF':cexf,'calificacion_final':final,'situacion':sit,'umbral_aprobacion':umbral}

@app.get('/api/anios-escolares')
def listar_anios_escolares(u=Depends(current_user)):
    return rows('SELECT * FROM anios_escolares ORDER BY activo DESC, nombre DESC')

@app.post('/api/anios-escolares')
def crear_anio_escolar(x:AnioEscolarIn,u=Depends(admin_user)):
    if x.fecha_inicio and x.fecha_fin and x.fecha_fin < x.fecha_inicio: raise HTTPException(400,'La fecha final no puede ser anterior a la fecha inicial')
    c=conn()
    try:
        cur=c.execute('INSERT INTO anios_escolares(nombre,fecha_inicio,fecha_fin,activo,cerrado,fecha_registro) VALUES(?,?,?,?,?,?)',(x.nombre.strip(),x.fecha_inicio.isoformat() if x.fecha_inicio else None,x.fecha_fin.isoformat() if x.fecha_fin else None,int(x.activo),int(x.cerrado),now()))
        aid=cur.lastrowid
        if x.activo:
            c.execute('UPDATE anios_escolares SET activo=0 WHERE id<>?',(aid,))
            c.execute("UPDATE configuracion SET valor=? WHERE clave='anio_escolar'",(x.nombre.strip(),))
        for n in range(1,5): c.execute('INSERT INTO periodos_academicos(anio_escolar_id,numero,nombre,estado,fecha_registro) VALUES(?,?,?,?,?)',(aid,n,f'Período {n}','Programado',now()))
        c.commit(); out=dict(c.execute('SELECT * FROM anios_escolares WHERE id=?',(aid,)).fetchone()); return out
    except sqlite3.IntegrityError: c.rollback(); raise HTTPException(400,'El año escolar ya existe')
    finally: c.close()

@app.get('/api/anios-escolares/actual')
def anio_escolar_actual(u=Depends(current_user)):
    r=rows('SELECT * FROM anios_escolares WHERE activo=1 LIMIT 1')
    return r[0] if r else None

@app.put('/api/anios-escolares/{anio_id}')
def actualizar_anio_escolar(anio_id:int,x:AnioEscolarIn,u=Depends(admin_user)):
    if x.fecha_inicio and x.fecha_fin and x.fecha_fin < x.fecha_inicio: raise HTTPException(400,'La fecha final no puede ser anterior a la fecha inicial')
    c=conn(); r=c.execute('SELECT * FROM anios_escolares WHERE id=?',(anio_id,)).fetchone()
    if not r: c.close(); raise HTTPException(404,'Año escolar no encontrado')
    if r['cerrado'] and not x.cerrado: c.close(); raise HTTPException(400,'Un año escolar cerrado no puede reabrirse desde este módulo')
    try:
        c.execute('UPDATE anios_escolares SET nombre=?,fecha_inicio=?,fecha_fin=?,cerrado=? WHERE id=?',(x.nombre.strip(),x.fecha_inicio.isoformat() if x.fecha_inicio else None,x.fecha_fin.isoformat() if x.fecha_fin else None,int(x.cerrado),anio_id))
        if x.activo:
            c.execute('UPDATE anios_escolares SET activo=0 WHERE id<>?',(anio_id,)); c.execute("UPDATE configuracion SET valor=? WHERE clave='anio_escolar'",(x.nombre.strip(),))
            c.execute('UPDATE anios_escolares SET activo=1 WHERE id=?',(anio_id,))
        c.commit(); return dict(c.execute('SELECT * FROM anios_escolares WHERE id=?',(anio_id,)).fetchone())
    except sqlite3.IntegrityError: c.rollback(); raise HTTPException(400,'El año escolar ya existe')
    finally: c.close()

@app.post('/api/anios-escolares/{anio_id}/activar')
def activar_anio_escolar(anio_id:int,u=Depends(admin_user)):
    c=conn(); r=c.execute('SELECT * FROM anios_escolares WHERE id=?',(anio_id,)).fetchone()
    if not r: c.close(); raise HTTPException(404,'Año escolar no encontrado')
    if r['cerrado']: c.close(); raise HTTPException(400,'No se puede activar un año escolar cerrado')
    c.execute('UPDATE anios_escolares SET activo=0'); c.execute('UPDATE anios_escolares SET activo=1 WHERE id=?',(anio_id,)); c.execute("UPDATE configuracion SET valor=? WHERE clave='anio_escolar'",(r['nombre'],)); c.commit(); out=dict(c.execute('SELECT * FROM anios_escolares WHERE id=?',(anio_id,)).fetchone()); c.close(); return out

@app.get('/api/periodos-academicos')
def listar_periodos_academicos(anio_escolar_id:Optional[int]=None,u=Depends(current_user)):
    if anio_escolar_id is None:
        a=rows('SELECT id FROM anios_escolares WHERE activo=1 LIMIT 1'); anio_escolar_id=a[0]['id'] if a else None
    if anio_escolar_id is None: return []
    return rows('SELECT p.*,a.nombre anio_nombre FROM periodos_academicos p JOIN anios_escolares a ON a.id=p.anio_escolar_id WHERE p.anio_escolar_id=? ORDER BY p.numero',(anio_escolar_id,))

@app.get('/api/periodos-academicos/actual')
def periodo_academico_actual(u=Depends(current_user)):
    a=rows('SELECT * FROM anios_escolares WHERE activo=1 LIMIT 1')
    if not a: return None
    aid=a[0]['id']; hoy=date.today().isoformat()
    r=rows("SELECT p.*,a.nombre anio_nombre FROM periodos_academicos p JOIN anios_escolares a ON a.id=p.anio_escolar_id WHERE p.anio_escolar_id=? AND ((p.fecha_inicio IS NOT NULL AND p.fecha_fin IS NOT NULL AND p.fecha_inicio<=? AND p.fecha_fin>=?) OR p.estado='Activo') ORDER BY CASE WHEN p.estado='Activo' THEN 0 ELSE 1 END,p.numero LIMIT 1",(aid,hoy,hoy))
    return r[0] if r else None

@app.post('/api/periodos-academicos')
def crear_periodo_academico(x:PeriodoAcademicoIn,u=Depends(admin_user)):
    if x.fecha_inicio and x.fecha_fin and x.fecha_fin < x.fecha_inicio: raise HTTPException(400,'La fecha final no puede ser anterior a la fecha inicial')
    if not rows('SELECT id FROM anios_escolares WHERE id=?',(x.anio_escolar_id,)): raise HTTPException(404,'Año escolar no encontrado')
    c=conn()
    try:
        cur=c.execute('INSERT INTO periodos_academicos(anio_escolar_id,numero,nombre,fecha_inicio,fecha_fin,estado,fecha_registro) VALUES(?,?,?,?,?,?,?)',(x.anio_escolar_id,x.numero,x.nombre.strip(),x.fecha_inicio.isoformat() if x.fecha_inicio else None,x.fecha_fin.isoformat() if x.fecha_fin else None,x.estado,now()))
        c.commit(); out=dict(c.execute('SELECT * FROM periodos_academicos WHERE id=?',(cur.lastrowid,)).fetchone()); return out
    except sqlite3.IntegrityError: c.rollback(); raise HTTPException(400,'Ya existe ese período para el año escolar')
    finally: c.close()

@app.put('/api/periodos-academicos/{periodo_id}')
def actualizar_periodo_academico(periodo_id:int,x:PeriodoAcademicoIn,u=Depends(admin_user)):
    if x.fecha_inicio and x.fecha_fin and x.fecha_fin < x.fecha_inicio: raise HTTPException(400,'La fecha final no puede ser anterior a la fecha inicial')
    c=conn(); r=c.execute('SELECT * FROM periodos_academicos WHERE id=?',(periodo_id,)).fetchone()
    if not r: c.close(); raise HTTPException(404,'Período académico no encontrado')
    try:
        c.execute('UPDATE periodos_academicos SET anio_escolar_id=?,numero=?,nombre=?,fecha_inicio=?,fecha_fin=?,estado=? WHERE id=?',(x.anio_escolar_id,x.numero,x.nombre.strip(),x.fecha_inicio.isoformat() if x.fecha_inicio else None,x.fecha_fin.isoformat() if x.fecha_fin else None,x.estado,periodo_id)); c.commit(); return dict(c.execute('SELECT * FROM periodos_academicos WHERE id=?',(periodo_id,)).fetchone())
    except sqlite3.IntegrityError: c.rollback(); raise HTTPException(400,'Ya existe ese período para el año escolar')
    finally: c.close()

@app.get('/api/auditoria')
def listar_auditoria(limit:int=100,usuario_id:Optional[int]=None,ruta:Optional[str]=None,u=Depends(admin_user)):
    limit=max(1,min(limit,500)); sql='SELECT * FROM auditoria WHERE 1=1'; args=[]
    if usuario_id is not None: sql+=' AND usuario_id=?'; args.append(usuario_id)
    if ruta: sql+=' AND ruta LIKE ?'; args.append('%'+ruta.strip()+'%')
    sql+=' ORDER BY id DESC LIMIT ?'; args.append(limit)
    return rows(sql,args)

@app.delete('/api/auditoria')
def limpiar_auditoria(u=Depends(admin_user)):
    c=conn(); c.execute('DELETE FROM auditoria'); c.commit(); c.close(); audit_log(u['id'],u['usuario'],'DELETE','/api/auditoria',200,None,'Limpieza manual de auditoría'); return {'ok':True}

@app.get('/api/institucion')
def obtener_institucion(u=Depends(current_user)):
    keys=['nombre_centro','codigo_centro','distrito','regional','direccion','telefono','correo_centro','director']
    d={r['clave']:r['valor'] for r in rows('SELECT clave,valor FROM configuracion WHERE clave IN ('+','.join('?'*len(keys))+')',keys)}
    return {'nombre_centro':d.get('nombre_centro',''),'codigo_centro':d.get('codigo_centro',''),'distrito':d.get('distrito',''),'regional':d.get('regional',''),'direccion':d.get('direccion',''),'telefono':d.get('telefono',''),'correo':d.get('correo_centro',''),'director':d.get('director','')}

@app.put('/api/institucion')
def actualizar_institucion(x:InstitucionUpdate,u=Depends(admin_user)):
    vals={'nombre_centro':x.nombre_centro.strip(),'codigo_centro':x.codigo_centro or '','distrito':x.distrito or '','regional':x.regional or '','direccion':x.direccion or '','telefono':x.telefono or '','correo_centro':x.correo or '','director':x.director or ''}
    c=conn();
    for k,v in vals.items(): c.execute('INSERT INTO configuracion(clave,valor) VALUES(?,?) ON CONFLICT(clave) DO UPDATE SET valor=excluded.valor',(k,v))
    c.commit(); c.close(); audit_log(u['id'],u['usuario'],'PUT','/api/institucion',200,None,'Actualización de datos institucionales'); return obtener_institucion(u)

@app.get('/api/configuracion')
def listar_configuracion(u=Depends(current_user)):
    return rows('SELECT clave,valor FROM configuracion ORDER BY clave')

@app.put('/api/configuracion/{clave}')
def actualizar_configuracion(clave:str,x:ConfigUpdate,u=Depends(admin_user)):
    allowed={'anio_escolar','max_calificacion','rp_reemplaza_p','final_promedio_competencias','nombre_centro','umbral_aprobacion','peso_completiva_cf','peso_completiva_cec','peso_extra_cf','peso_extra_ceex','codigo_centro','distrito','regional','direccion','telefono','correo_centro','director'}
    if clave not in allowed: raise HTTPException(400,'Clave de configuración no permitida')
    c=conn(); c.execute('INSERT INTO configuracion(clave,valor) VALUES(?,?) ON CONFLICT(clave) DO UPDATE SET valor=excluded.valor',(clave,x.valor)); c.commit(); c.close(); return {'clave':clave,'valor':x.valor}

@app.get('/api/cierre/asignatura')
def cierre_asignatura(curso_id:int,asignatura_id:int,u=Depends(current_user)):
    if not can_access_course(u,curso_id): raise HTTPException(403,'No tienes acceso a esta clase')
    if not rows('SELECT id FROM curso_asignatura WHERE curso_id=? AND asignatura_id=? AND activa=1',(curso_id,asignatura_id)):
        raise HTTPException(400,'La asignatura no está asociada a la clase')
    students=rows("SELECT id,matricula,nombre FROM estudiantes WHERE curso_id=? AND estado='Activo' ORDER BY nombre COLLATE NOCASE,id",(curso_id,))
    out=[]
    for st in students:
        sid=st['id']; cf=_calificacion_final_base(sid,asignatura_id)
        saved=rows('SELECT completiva,extraordinaria,especial_cf,especial_ce,situacion,observacion FROM calificaciones_finales WHERE estudiante_id=? AND asignatura_id=?',(sid,asignatura_id))
        r=saved[0] if saved else {'completiva':None,'extraordinaria':None,'especial_cf':None,'especial_ce':None,'situacion':None,'observacion':None}
        rr=_resultado_final(cf,r); comps=[]
        for c in rows('SELECT id,codigo,nombre FROM competencias WHERE asignatura_id=? AND activa=1 ORDER BY orden,codigo',(asignatura_id,)):
            pv=[_periodo_competencia(sid,asignatura_id,c['id'],p) for p in range(1,5)]; rv=[_rp_competencia(sid,asignatura_id,c['id'],p) for p in range(1,5)]; vals=[]
            for p in range(4):
                if _cfg('rp_reemplaza_p','1')=='1' and rv[p] is not None: vals.append(rv[p])
                elif pv[p] is not None: vals.append(pv[p])
                elif rv[p] is not None: vals.append(rv[p])
            comps.append({'id':c['id'],'codigo':c['codigo'],'nombre':c['nombre'],'P1':pv[0],'P2':pv[1],'P3':pv[2],'P4':pv[3],'RP1':rv[0],'RP2':rv[1],'RP3':rv[2],'RP4':rv[3],'promedio_anual':round(sum(vals)/len(vals),2) if vals else None})
        out.append({'estudiante_id':sid,'matricula':st['matricula'],'estudiante':st['nombre'],'competencias':comps,**r,**rr})
    return {'curso_id':curso_id,'asignatura_id':asignatura_id,'config':{k:_cfg(k) for k in ['anio_escolar','rp_reemplaza_p','umbral_aprobacion','peso_completiva_cf','peso_completiva_cec','peso_extra_cf','peso_extra_ceex']},'filas':out}

@app.get('/api/cierre/estudiante/{estudiante_id}')
def cierre_estudiante(estudiante_id:int,u=Depends(current_user)):
    if not can_access_student(u,estudiante_id): raise HTTPException(403,'No tienes acceso a este estudiante')
    st=rows('SELECT id,matricula,nombre,curso_id FROM estudiantes WHERE id=?',(estudiante_id,))
    if not st: raise HTTPException(404,'Estudiante no encontrado')
    subs=rows('SELECT a.id,a.nombre,a.codigo FROM asignaturas a JOIN curso_asignatura ca ON ca.asignatura_id=a.id WHERE ca.curso_id=? AND ca.activa=1 ORDER BY a.nombre',(st[0]['curso_id'],)); out=[]
    for a in subs:
        cf=_calificacion_final_base(estudiante_id,a['id']); saved=rows('SELECT completiva,extraordinaria,especial_cf,especial_ce,situacion,observacion FROM calificaciones_finales WHERE estudiante_id=? AND asignatura_id=?',(estudiante_id,a['id']))
        r=saved[0] if saved else {'completiva':None,'extraordinaria':None,'especial_cf':None,'especial_ce':None,'situacion':None,'observacion':None}
        out.append({'asignatura_id':a['id'],'asignatura':a['nombre'],'codigo':a['codigo'],**r,**_resultado_final(cf,r)})
    return {'estudiante':st[0],'filas':out}

@app.post('/api/calificaciones-finales')
def guardar_final(x:CalificacionFinalIn,u=Depends(current_user)):
    if not can_access_student(u,x.estudiante_id): raise HTTPException(403,'No tienes acceso a este estudiante')
    if not rows('SELECT id FROM asignaturas WHERE id=? AND activa=1',(x.asignatura_id,)): raise HTTPException(404,'Asignatura no encontrada')
    r=x.model_dump(); cf=_calificacion_final_base(x.estudiante_id,x.asignatura_id); result=_resultado_final(cf,r)
    c=conn(); c.execute('''INSERT INTO calificaciones_finales(estudiante_id,asignatura_id,completiva,extraordinaria,especial_cf,especial_ce,situacion,observacion,fecha_actualizacion) VALUES(?,?,?,?,?,?,?,?,?) ON CONFLICT(estudiante_id,asignatura_id) DO UPDATE SET completiva=excluded.completiva,extraordinaria=excluded.extraordinaria,especial_cf=excluded.especial_cf,especial_ce=excluded.especial_ce,situacion=excluded.situacion,observacion=excluded.observacion,fecha_actualizacion=excluded.fecha_actualizacion''',(x.estudiante_id,x.asignatura_id,x.completiva,x.extraordinaria,x.especial_cf,x.especial_ce,result['situacion'],x.observacion,now())); c.commit(); c.close(); return {'ok':True,**result}



# ============================================================
# ETAPA 8 — DASHBOARD, INDICADORES, ALERTAS Y BÚSQUEDA GLOBAL
# ============================================================
def _scope_course_ids(u):
    return allowed_course_ids(u)

def _in_scope_course(u, cid):
    return cid is not None and can_access_course(u, int(cid))

@app.get('/api/dashboard')
def dashboard(curso_id:Optional[int]=None, fecha:Optional[date]=None, u=Depends(current_user)):
    target_date=(fecha or date.today()).isoformat()
    ids=_scope_course_ids(u)
    if curso_id is not None and not _in_scope_course(u,curso_id):
        raise HTTPException(403,'No tienes acceso a esta clase')
    course_clause=''; args=[]
    if u['rol']!='admin':
        if not ids: return {'fecha':target_date,'resumen':{},'asistencia':{},'academico':{},'alertas':[],'cursos':[]}
        course_clause=' AND e.curso_id IN ('+','.join('?'*len(ids))+')'; args.extend(ids)
    if curso_id is not None:
        course_clause+=' AND e.curso_id=?'; args.append(curso_id)
    students=rows('SELECT e.id,e.curso_id,e.nombre,e.matricula,c.nombre curso_nombre FROM estudiantes e LEFT JOIN cursos c ON c.id=e.curso_id WHERE e.estado=\'Activo\''+course_clause+' ORDER BY e.nombre COLLATE NOCASE',args)
    student_ids=[x['id'] for x in students]
    base_course=''
    cargs=[]
    if u['rol']!='admin':
        base_course=' WHERE c.id IN ('+','.join('?'*len(ids))+')'; cargs.extend(ids)
    if curso_id is not None:
        base_course=(' WHERE ' if not base_course else base_course+' AND ')+'c.id=?'; cargs.append(curso_id)
    total_students=len(students)
    courses=rows('SELECT c.id,c.nombre,c.grado,c.seccion,d.nombre docente_nombre FROM cursos c LEFT JOIN docentes d ON d.id=c.docente_id'+base_course+' ORDER BY c.nombre COLLATE NOCASE',cargs)
    qargs=[]
    q='SELECT COUNT(*) n FROM curso_asignatura ca JOIN cursos c ON c.id=ca.curso_id WHERE ca.activa=1'
    if u['rol']!='admin': q+=' AND c.id IN ('+','.join('?'*len(ids))+')'; qargs.extend(ids)
    if curso_id is not None: q+=' AND c.id=?'; qargs.append(curso_id)
    subject_count=rows(q,qargs)[0]['n'] if rows(q,qargs) else 0
    teacher_q='SELECT COUNT(*) n FROM docentes WHERE estado=\'Activo\''; teacher_count=rows(teacher_q)[0]['n'] if rows(teacher_q) else 0
    if student_ids:
        marks=','.join('?'*len(student_ids)); att=rows(f'''SELECT COUNT(*) total,
          SUM(CASE WHEN estado='Presente' THEN 1 ELSE 0 END) presentes,
          SUM(CASE WHEN estado='Ausente' THEN 1 ELSE 0 END) ausentes,
          SUM(CASE WHEN estado='Tardanza' THEN 1 ELSE 0 END) tardanzas,
          SUM(CASE WHEN estado='Excusa' THEN 1 ELSE 0 END) excusas
          FROM asistencia WHERE fecha=? AND estudiante_id IN ({marks})''',[target_date,*student_ids])[0]
        att_rows=rows(f'''SELECT e.id,e.nombre,e.matricula,e.curso_id,
          COUNT(a.id) registros,
          SUM(CASE WHEN a.estado IN ('Presente','Tardanza','Excusa') THEN 1 ELSE 0 END) favorables,
          SUM(CASE WHEN a.estado='Ausente' THEN 1 ELSE 0 END) ausencias
          FROM estudiantes e LEFT JOIN asistencia a ON a.estudiante_id=e.id
          WHERE e.estado='Activo' AND e.id IN ({marks}) GROUP BY e.id ORDER BY e.nombre COLLATE NOCASE''',student_ids)
    else:
        att={'total':0,'presentes':0,'ausentes':0,'tardanzas':0,'excusas':0}; att_rows=[]
    attendance_rate=round((att['presentes'] or 0)*100/(att['total'] or 1),2) if att['total'] else 0
    # Rendimiento académico: promedio de C.F. disponibles dentro del alcance.
    acad=[]
    for st in students:
        vals=[]
        for a in rows('SELECT DISTINCT ca.asignatura_id FROM curso_asignatura ca WHERE ca.curso_id=? AND ca.activa=1',(st['curso_id'],)):
            cf=_calificacion_final_base(st['id'],a['asignatura_id'])
            if cf is not None: vals.append(cf)
        if vals: acad.append({'id':st['id'],'nombre':st['nombre'],'matricula':st['matricula'],'curso_id':st['curso_id'],'promedio':round(sum(vals)/len(vals),2)})
    academic_avg=round(sum(x['promedio'] for x in acad)/len(acad),2) if acad else 0
    low_att=[]
    for x in att_rows:
        if x['registros'] and (x['favorables']*100/x['registros'])<80:
            low_att.append({'tipo':'Asistencia','estudiante_id':x['id'],'estudiante':x['nombre'],'detalle':f"{round(x['favorables']*100/x['registros'],1)}% de registros favorables"})
    low_perf=[{'tipo':'Rendimiento','estudiante_id':x['id'],'estudiante':x['nombre'],'detalle':f"Promedio {x['promedio']}"} for x in acad if x['promedio']<70]
    pending=[]; incidents=[]
    if student_ids:
        marks=','.join('?'*len(student_ids))
        pending=rows(f'''SELECT p.id,e.nombre estudiante,p.fecha,p.tipo,p.motivo FROM permisos p JOIN estudiantes e ON e.id=p.estudiante_id WHERE p.estado='Pendiente' AND p.estudiante_id IN ({marks}) ORDER BY p.fecha DESC''',student_ids)
        incidents=rows(f'''SELECT i.id,e.nombre estudiante,i.fecha,i.tipo,i.gravedad,i.descripcion FROM incidencias i JOIN estudiantes e ON e.id=i.estudiante_id WHERE i.estudiante_id IN ({marks}) AND i.fecha>=date(?, '-30 day') ORDER BY i.fecha DESC LIMIT 20''',[*student_ids,target_date])
    alerts=low_att+low_perf+[{'tipo':'Permiso pendiente','estudiante':x['estudiante'],'detalle':f"{x['fecha']} · {x['tipo']}"} for x in pending]
    return {'fecha':target_date,'resumen':{'clases':len(courses),'estudiantes':total_students,'asignaturas':subject_count,'docentes':teacher_count},'asistencia':{'registros':att['total'] or 0,'presentes':att['presentes'] or 0,'ausentes':att['ausentes'] or 0,'tardanzas':att['tardanzas'] or 0,'excusas':att['excusas'] or 0,'porcentaje':attendance_rate},'academico':{'promedio_general':academic_avg,'estudiantes_con_datos':len(acad),'bajo_umbral':len(low_perf)},'alertas':alerts[:50],'incidencias_30_dias':incidents,'cursos':courses}

@app.get('/api/busqueda')
def busqueda(q:str, u=Depends(current_user)):
    q=(q or '').strip()
    if len(q)<2: return {'estudiantes':[],'cursos':[],'asignaturas':[],'docentes':[]}
    like='%'+q+'%'; ids=allowed_course_ids(u)
    if u['rol']=='admin':
        est=rows('SELECT id,nombre,matricula,curso_id FROM estudiantes WHERE estado=\'Activo\' AND (nombre LIKE ? OR matricula LIKE ?) ORDER BY nombre LIMIT 20',(like,like))
        cur=rows('SELECT id,nombre,grado,seccion FROM cursos WHERE nombre LIKE ? OR grado LIKE ? OR seccion LIKE ? ORDER BY nombre LIMIT 20',(like,like,like))
        asi=rows('SELECT id,nombre,codigo,grado FROM asignaturas WHERE activa=1 AND (nombre LIKE ? OR codigo LIKE ?) ORDER BY nombre LIMIT 20',(like,like))
        doc=rows('SELECT id,nombre,especialidad FROM docentes WHERE nombre LIKE ? OR especialidad LIKE ? ORDER BY nombre LIMIT 20',(like,like))
    else:
        if not ids:return {'estudiantes':[],'cursos':[],'asignaturas':[],'docentes':[]}
        marks=','.join('?'*len(ids))
        est=rows(f'SELECT id,nombre,matricula,curso_id FROM estudiantes WHERE estado=\'Activo\' AND curso_id IN ({marks}) AND (nombre LIKE ? OR matricula LIKE ?) ORDER BY nombre LIMIT 20',[*ids,like,like])
        cur=rows(f'SELECT id,nombre,grado,seccion FROM cursos WHERE id IN ({marks}) AND (nombre LIKE ? OR grado LIKE ? OR seccion LIKE ?) ORDER BY nombre LIMIT 20',[*ids,like,like,like])
        pairs=allowed_subject_pairs(u); aids=sorted(set(x['asignatura_id'] for x in pairs)); asi=rows('SELECT id,nombre,codigo,grado FROM asignaturas WHERE activa=1 AND id IN ('+','.join('?'*len(aids))+') AND (nombre LIKE ? OR codigo LIKE ?) ORDER BY nombre LIMIT 20',[*aids,like,like]) if aids else []
        did=u.get('docente_id'); doc=rows('SELECT id,nombre,especialidad FROM docentes WHERE id=? AND (nombre LIKE ? OR especialidad LIKE ?)',(did,like,like)) if did else []
    return {'estudiantes':est,'cursos':cur,'asignaturas':asi,'docentes':doc}

# ============================================================
# ETAPA 8 — REPORTES Y DOCUMENTOS
# ============================================================
def _safe_filename(text):
    return re.sub(r'[^A-Za-z0-9._-]+','_',str(text or 'reporte')).strip('_') or 'reporte'

def _course_or_403(u, curso_id):
    if not can_access_course(u, curso_id): raise HTTPException(403,'No tienes acceso a esta clase')
    c=rows('SELECT id,nombre,grado,seccion,anio_escolar FROM cursos WHERE id=?',(curso_id,))
    if not c: raise HTTPException(404,'Clase no encontrada')
    return c[0]

def _report_grade_rows(curso_id, asignatura_id, u):
    _course_or_403(u,curso_id)
    if not rows('SELECT id FROM curso_asignatura WHERE curso_id=? AND asignatura_id=? AND activa=1',(curso_id,asignatura_id)):
        raise HTTPException(400,'La asignatura no está asociada a la clase')
    return cierre_asignatura(curso_id,asignatura_id,u)

def _workbook_bytes(wb):
    out=io.BytesIO(); wb.save(out); out.seek(0); return out

def _xlsx_bytes(headers, rows_data, sheet='Reporte'):
    wb=Workbook(); ws=wb.active; ws.title=sheet[:31]
    ws.append(headers)
    for row in rows_data: ws.append(list(row))
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:
        maxlen=max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width=min(max(maxlen+2,12),40)
    out=io.BytesIO(); wb.save(out); out.seek(0); return out

def _pdf_bytes(title, headers, rows_data, subtitle=None):
    out=io.BytesIO(); doc=SimpleDocTemplate(out,pagesize=landscape(letter),rightMargin=24,leftMargin=24,topMargin=24,bottomMargin=24)
    styles=getSampleStyleSheet(); elems=[Paragraph(title,styles['Title'])]
    if subtitle: elems += [Paragraph(subtitle,styles['Normal']),Spacer(1,10)]
    clean=[[str(x if x is not None else '—') for x in r] for r in rows_data]
    table=Table([headers]+clean,repeatRows=1)
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e5e7eb')),('GRID',(0,0),(-1,-1),0.4,colors.grey),('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('FONTSIZE',(0,0),(-1,-1),7),('VALIGN',(0,0),(-1,-1),'MIDDLE')]))
    elems.append(table); doc.build(elems); out.seek(0); return out

def _attendance_report(curso_id, fecha_desde=None, fecha_hasta=None, u=None):
    _course_or_403(u,curso_id)
    q="""SELECT e.matricula,e.nombre,
      SUM(CASE WHEN a.estado='Presente' THEN 1 ELSE 0 END) presentes,
      SUM(CASE WHEN a.estado='Ausente' THEN 1 ELSE 0 END) ausentes,
      SUM(CASE WHEN a.estado='Tardanza' THEN 1 ELSE 0 END) tardanzas,
      SUM(CASE WHEN a.estado='Excusa' THEN 1 ELSE 0 END) excusas,
      COUNT(a.id) registros
      FROM estudiantes e LEFT JOIN asistencia a ON a.estudiante_id=e.id"""
    args=[curso_id]; clauses=["e.curso_id=?","e.estado='Activo'"]
    if fecha_desde: clauses.append('a.fecha>=?'); args.append(fecha_desde.isoformat())
    if fecha_hasta: clauses.append('a.fecha<=?'); args.append(fecha_hasta.isoformat())
    q+=' WHERE '+' AND '.join(clauses)+' GROUP BY e.id ORDER BY e.nombre COLLATE NOCASE,e.id'
    return rows(q,args)

# ===================== ETAPA 19: ADMINISTRACION AVANZADA =====================

def _csv_rows(text):
    text=text.lstrip('\ufeff').strip()
    if not text: return []
    try: dialect=csv.Sniffer().sniff(text[:4096], delimiters=',;\t')
    except Exception: dialect=csv.excel
    return list(csv.DictReader(io.StringIO(text), dialect=dialect))

@app.post('/api/admin/importar/docentes')
def importar_docentes(x:ImportarDocentesIn,u=Depends(admin_user)):
    data=_csv_rows(x.csv_text)
    if not data: raise HTTPException(400,'El CSV no contiene registros')
    c=conn(); creados=actualizados=0; errores=[]
    for n,row in enumerate(data,2):
        nombre=(row.get('nombre') or row.get('Nombre') or '').strip()
        if not nombre: errores.append({'fila':n,'error':'Falta nombre'}); continue
        cedula=(row.get('cedula') or row.get('Cédula') or '').strip() or None; correo=(row.get('correo') or row.get('Correo') or '').strip() or None; telefono=(row.get('telefono') or row.get('Teléfono') or '').strip() or None; especialidad=(row.get('especialidad') or row.get('Especialidad') or row.get('area') or '').strip() or None; estado=(row.get('estado') or row.get('Estado') or 'Activo').strip() or 'Activo'
        existe=c.execute('SELECT id FROM docentes WHERE cedula=?',(cedula,)).fetchone() if cedula else None
        if existe: c.execute('UPDATE docentes SET nombre=?,correo=?,telefono=?,especialidad=?,estado=? WHERE id=?',(nombre,correo,telefono,especialidad,estado,existe['id'])); actualizados+=1
        else: c.execute('INSERT INTO docentes(nombre,cedula,correo,telefono,especialidad,estado,fecha_registro) VALUES(?,?,?,?,?,?,?)',(nombre,cedula,correo,telefono,especialidad,estado,now())); creados+=1
    c.commit(); c.close(); audit_log(u['id'],u['usuario'],'POST','/api/admin/importar/docentes',200,None,f'Importación docentes: {creados} creados, {actualizados} actualizados')
    return {'ok':True,'creados':creados,'actualizados':actualizados,'errores':errores}

@app.post('/api/admin/importar/estudiantes')
def importar_estudiantes(x:ImportarEstudiantesIn,u=Depends(admin_user)):
    data=_csv_rows(x.csv_text)
    if not data: raise HTTPException(400,'El CSV no contiene registros')
    c=conn(); creados=actualizados=0; errores=[]
    for n,row in enumerate(data,2):
        nombres=(row.get('nombres') or row.get('Nombres') or '').strip(); apellidos=(row.get('apellidos') or row.get('Apellidos') or '').strip(); fn=(row.get('fecha_nacimiento') or row.get('Fecha nacimiento') or row.get('fecha') or '').strip(); cid=(row.get('curso_id') or row.get('Curso ID') or '').strip(); curso_nombre=(row.get('curso') or row.get('Curso') or '').strip()
        if not nombres or not apellidos or not fn: errores.append({'fila':n,'error':'Se requieren nombres, apellidos y fecha_nacimiento'}); continue
        if not cid and curso_nombre:
            rr=c.execute('SELECT id FROM cursos WHERE nombre=?',(curso_nombre,)).fetchone(); cid=str(rr['id']) if rr else ''
        try: cid=int(cid)
        except: errores.append({'fila':n,'error':'curso_id no válido o curso no encontrado'}); continue
        if not c.execute('SELECT id FROM cursos WHERE id=?',(cid,)).fetchone(): errores.append({'fila':n,'error':'Clase inexistente'}); continue
        try: date.fromisoformat(fn)
        except: errores.append({'fila':n,'error':'fecha_nacimiento debe tener formato YYYY-MM-DD'}); continue
        correo=(row.get('correo') or row.get('Correo') or '').strip() or None; telefono=(row.get('telefono') or row.get('Teléfono') or '').strip() or None; estado=(row.get('estado') or row.get('Estado') or 'Activo').strip() or 'Activo'; nombre_completo=f'{nombres} {apellidos}'.strip(); matricula=(row.get('matricula') or row.get('Matrícula') or '').strip() or None
        existe=c.execute('SELECT * FROM estudiantes WHERE matricula=?',(matricula,)).fetchone() if matricula else c.execute('SELECT * FROM estudiantes WHERE nombre=? AND fecha_nacimiento=? AND curso_id=?',(nombre_completo,fn,cid)).fetchone()
        if existe:
            c.execute('UPDATE estudiantes SET nombre=?,curso_id=?,correo=?,telefono=?,estado=?,fecha_nacimiento=? WHERE id=?',(nombre_completo,cid,correo,telefono,estado,fn,existe['id'])); actualizados+=1
        else:
            initials=(initial(nombres)+initial(apellidos)).upper(); orden=int(c.execute('SELECT COALESCE(MAX(orden_lista),0) m FROM estudiantes WHERE curso_id=?',(cid,)).fetchone()['m'] or 0)+1; mat=matricula or f'{initials}{fn[:4]}-{orden}'
            while c.execute('SELECT 1 FROM estudiantes WHERE matricula=?',(mat,)).fetchone(): orden+=1; mat=f'{initials}{fn[:4]}-{orden}'
            curso_row=c.execute('SELECT grado,seccion FROM cursos WHERE id=?',(cid,)).fetchone(); cur=c.execute('INSERT INTO estudiantes(matricula,nombre,curso,seccion,correo,telefono,estado,fecha_registro,fecha_nacimiento,edad_registro,curso_id,orden_lista) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',(mat,nombre_completo,curso_row['grado'] if curso_row else None,curso_row['seccion'] if curso_row else None,correo,telefono,estado,now(),fn,age(date.fromisoformat(fn)),cid,orden)); sid=cur.lastrowid; c.execute('INSERT OR IGNORE INTO estudiante_curso(estudiante_id,curso_id,fecha_inicio) VALUES(?,?,?)',(sid,cid,date.today().isoformat())); creados+=1
    c.commit(); c.close(); audit_log(u['id'],u['usuario'],'POST','/api/admin/importar/estudiantes',200,None,f'Importación estudiantes: {creados} creados, {actualizados} actualizados')
    return {'ok':True,'creados':creados,'actualizados':actualizados,'errores':errores}

@app.get('/api/admin/exportar/general.xlsx')
def exportar_general(u=Depends(admin_user)):
    wb=Workbook(); wb.remove(wb.active)
    tablas=['usuarios','docentes','cursos','asignaturas','curso_asignatura','estudiantes','estudiante_curso','competencias','actividades','calificaciones','asistencia','incidencias','permisos','eventos_calendario','horario_escolar','anios_escolares','periodos_academicos','planificaciones','seguimientos','reuniones_familia','compromisos','alertas_riesgo','planes_intervencion','plan_acciones','comunicados','comunicado_destinatarios','calificaciones_finales','configuracion']
    for tabla in tablas:
        ws=wb.create_sheet(tabla[:31]); rs=rows(f'SELECT * FROM {tabla}')
        if rs:
            headers=list(rs[0].keys()); ws.append(headers)
            for r in rs: ws.append([r.get(h) for h in headers])
        else: ws.append(['Sin registros'])
    data=_workbook_bytes(wb); audit_log(u['id'],u['usuario'],'GET','/api/admin/exportar/general.xlsx',200,None,'Exportación general XLSX')
    return StreamingResponse(data,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="gestion_clase_respaldo_general.xlsx"'})

@app.get('/api/admin/backup')
def descargar_backup(u=Depends(admin_user)):
    c=conn(); c.execute('PRAGMA wal_checkpoint(FULL)'); c.close(); data=Path(DB_PATH).read_bytes(); audit_log(u['id'],u['usuario'],'GET','/api/admin/backup',200,None,'Descarga de respaldo SQLite')
    return Response(content=data,media_type='application/x-sqlite3',headers={'Content-Disposition':'attachment; filename="gestion_clase_v5_1_backup.db"'})

@app.post('/api/admin/restore')
async def restaurar_backup(request:Request,u=Depends(admin_user)):
    raw=await request.body()
    if len(raw)<100: raise HTTPException(400,'Archivo de respaldo vacío o inválido')
    tmp=DB_PATH.with_suffix('.restore.tmp.db'); tmp.write_bytes(raw)
    try:
        tc=sqlite3.connect(tmp); ok=tc.execute('PRAGMA integrity_check').fetchone()[0]=='ok'; names={r[0] for r in tc.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()}; tc.close()
        if not ok or not {'usuarios','docentes','cursos','asignaturas','estudiantes'}.issubset(names): raise ValueError('La copia no corresponde a una base compatible')
        conn().close(); os.replace(tmp,DB_PATH); return {'ok':True,'mensaje':'Base restaurada. Reinicia el servidor para garantizar que las conexiones usen la copia restaurada.'}
    except Exception as e:
        try: tmp.unlink()
        except: pass
        raise HTTPException(400,f'Respaldo inválido: {e}')

@app.get('/api/reportes/cierre.csv')
def reporte_cierre_csv(curso_id:int,asignatura_id:int,u=Depends(current_user)):
    data=_report_grade_rows(curso_id,asignatura_id,u); output=io.StringIO(); w=csv.writer(output)
    headers=['Matrícula','Estudiante','C.F.','C.E.C.','C.C.F.','C.E.EX.','C.EX.F.','Especial C.F.','Especial C.E.','Calificación final','Situación']
    w.writerow(headers)
    for r in data['filas']: w.writerow([r['matricula'],r['estudiante'],r['CF'],r['completiva'],r['CCF'],r['extraordinaria'],r['CEXF'],r['especial_cf'],r['especial_ce'],r['calificacion_final'],r['situacion']])
    return StreamingResponse(iter([output.getvalue().encode('utf-8-sig')]),media_type='text/csv',headers={'Content-Disposition':f'attachment; filename="cierre_{curso_id}_{asignatura_id}.csv"'})

@app.get('/api/reportes/cierre.xlsx')
def reporte_cierre_xlsx(curso_id:int,asignatura_id:int,u=Depends(current_user)):
    data=_report_grade_rows(curso_id,asignatura_id,u); headers=['Matrícula','Estudiante','C.F.','C.E.C.','C.C.F.','C.E.EX.','C.EX.F.','Especial C.F.','Especial C.E.','Calificación final','Situación','Observación']
    vals=[[r['matricula'],r['estudiante'],r['CF'],r['completiva'],r['CCF'],r['extraordinaria'],r['CEXF'],r['especial_cf'],r['especial_ce'],r['calificacion_final'],r['situacion'],r.get('observacion','')] for r in data['filas']]
    b=_xlsx_bytes(headers,vals,'Cierre académico')
    return StreamingResponse(b,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="cierre_academico_{curso_id}_{asignatura_id}.xlsx"'})

@app.get('/api/reportes/cierre.pdf')
def reporte_cierre_pdf(curso_id:int,asignatura_id:int,u=Depends(current_user)):
    data=_report_grade_rows(curso_id,asignatura_id,u); course=_course_or_403(u,curso_id)
    headers=['Matrícula','Estudiante','C.F.','C.C.F.','C.EX.F.','Final','Situación']; vals=[[r['matricula'],r['estudiante'],r['CF'],r['CCF'],r['CEXF'],r['calificacion_final'],r['situacion']] for r in data['filas']]
    b=_pdf_bytes('Cierre académico',headers,vals,f"{course['nombre']} · {course.get('grado') or ''} {course.get('seccion') or ''} · Año {data['config']['anio_escolar']}")
    return StreamingResponse(b,media_type='application/pdf',headers={'Content-Disposition':f'attachment; filename="cierre_academico_{curso_id}_{asignatura_id}.pdf"'})

@app.get('/api/reportes/boletin/{estudiante_id}.xlsx')
def boletin_xlsx(estudiante_id:int,u=Depends(current_user)):
    if not can_access_student(u,estudiante_id): raise HTTPException(403,'No tienes acceso a este estudiante')
    data=cierre_estudiante(estudiante_id,u); st=data['estudiante']; headers=['Asignatura','Código','C.F.','C.E.C.','C.C.F.','C.E.EX.','C.EX.F.','Especial C.F.','Especial C.E.','Calificación final','Situación']
    vals=[[r['asignatura'],r['codigo'],r['CF'],r['completiva'],r['CCF'],r['extraordinaria'],r['CEXF'],r['especial_cf'],r['especial_ce'],r['calificacion_final'],r['situacion']] for r in data['filas']]
    b=_xlsx_bytes(headers,vals,'Boletín')
    return StreamingResponse(b,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="boletin_{_safe_filename(st["matricula"])}.xlsx"'})

@app.get('/api/reportes/boletin/{estudiante_id}.pdf')
def boletin_pdf(estudiante_id:int,u=Depends(current_user)):
    if not can_access_student(u,estudiante_id): raise HTTPException(403,'No tienes acceso a este estudiante')
    data=cierre_estudiante(estudiante_id,u); st=data['estudiante']; headers=['Asignatura','Código','C.F.','C.C.F.','C.EX.F.','Final','Situación']
    vals=[[r['asignatura'],r['codigo'],r['CF'],r['CCF'],r['CEXF'],r['calificacion_final'],r['situacion']] for r in data['filas']]
    b=_pdf_bytes('Boletín académico',headers,vals,f"{st['nombre']} · Matrícula {st['matricula']}")
    return StreamingResponse(b,media_type='application/pdf',headers={'Content-Disposition':f'attachment; filename="boletin_{_safe_filename(st["matricula"])}.pdf"'})

@app.get('/api/reportes/asistencia.xlsx')
def asistencia_xlsx(curso_id:int,fecha_desde:Optional[date]=None,fecha_hasta:Optional[date]=None,u=Depends(current_user)):
    data=_attendance_report(curso_id,fecha_desde,fecha_hasta,u); headers=['Matrícula','Estudiante','Presentes','Ausentes','Tardanzas','Excusas','Registros']; vals=[[r['matricula'],r['nombre'],r['presentes'],r['ausentes'],r['tardanzas'],r['excusas'],r['registros']] for r in data]
    b=_xlsx_bytes(headers,vals,'Asistencia')
    return StreamingResponse(b,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="asistencia_{curso_id}.xlsx"'})

@app.get('/api/reportes/asistencia.pdf')
def asistencia_pdf(curso_id:int,fecha_desde:Optional[date]=None,fecha_hasta:Optional[date]=None,u=Depends(current_user)):
    data=_attendance_report(curso_id,fecha_desde,fecha_hasta,u); c=_course_or_403(u,curso_id); headers=['Matrícula','Estudiante','Presentes','Ausentes','Tardanzas','Excusas','Registros']; vals=[[r['matricula'],r['nombre'],r['presentes'],r['ausentes'],r['tardanzas'],r['excusas'],r['registros']] for r in data]
    b=_pdf_bytes('Reporte de asistencia',headers,vals,f"{c['nombre']} · {fecha_desde or 'Inicio'} a {fecha_hasta or 'Actualidad'}")
    return StreamingResponse(b,media_type='application/pdf',headers={'Content-Disposition':f'attachment; filename="asistencia_{curso_id}.pdf"'})

@app.get('/api/reportes/incidencias.xlsx')
def incidencias_xlsx(curso_id:int,u=Depends(current_user)):
    _course_or_403(u,curso_id); data=rows("""SELECT e.matricula,e.nombre,i.fecha,i.hora,i.tipo,i.gravedad,i.descripcion,i.accion_tomada,i.observacion FROM incidencias i JOIN estudiantes e ON e.id=i.estudiante_id WHERE e.curso_id=? ORDER BY i.fecha DESC,e.nombre COLLATE NOCASE""",(curso_id,))
    headers=['Matrícula','Estudiante','Fecha','Hora','Tipo','Gravedad','Descripción','Acción tomada','Observación']; vals=[[r['matricula'],r['nombre'],r['fecha'],r['hora'],r['tipo'],r['gravedad'],r['descripcion'],r['accion_tomada'],r['observacion']] for r in data]
    b=_xlsx_bytes(headers,vals,'Incidencias'); return StreamingResponse(b,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="incidencias_{curso_id}.xlsx"'})

@app.get('/api/reportes/permisos.xlsx')
def permisos_xlsx(curso_id:int,u=Depends(current_user)):
    _course_or_403(u,curso_id); data=rows("""SELECT e.matricula,e.nombre,p.fecha,p.tipo,p.hora_desde,p.hora_hasta,p.motivo,p.estado,p.observacion FROM permisos p JOIN estudiantes e ON e.id=p.estudiante_id WHERE e.curso_id=? ORDER BY p.fecha DESC,e.nombre COLLATE NOCASE""",(curso_id,))
    headers=['Matrícula','Estudiante','Fecha','Tipo','Desde','Hasta','Motivo','Estado','Observación']; vals=[[r['matricula'],r['nombre'],r['fecha'],r['tipo'],r['hora_desde'],r['hora_hasta'],r['motivo'],r['estado'],r['observacion']] for r in data]
    b=_xlsx_bytes(headers,vals,'Permisos'); return StreamingResponse(b,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="permisos_{curso_id}.xlsx"'})

# ETAPA 11

def _risk_students(u):
    ids=allowed_course_ids(u)
    if u['rol']=='admin': return rows("SELECT e.id,e.nombre,e.matricula,e.curso_id FROM estudiantes e WHERE e.estado='Activo' ORDER BY e.nombre COLLATE NOCASE")
    if not ids:return []
    ph=','.join('?'*len(ids)); return rows(f"SELECT e.id,e.nombre,e.matricula,e.curso_id FROM estudiantes e WHERE e.estado='Activo' AND e.curso_id IN ({ph}) ORDER BY e.nombre COLLATE NOCASE",ids)

def _risk_metrics(sid):
    a=rows("SELECT COUNT(*) total,SUM(CASE WHEN estado='Ausente' THEN 1 ELSE 0 END) aus FROM asistencia WHERE estudiante_id=?",(sid,))[0]
    total=a['total'] or 0; aus=a['aus'] or 0; att=round((total-aus)*100/total,2) if total else None
    inc=rows("SELECT COUNT(*) n FROM incidencias WHERE estudiante_id=? AND fecha>=date('now','-30 day')",(sid,))[0]['n']
    grave=rows("SELECT COUNT(*) n FROM incidencias WHERE estudiante_id=? AND fecha>=date('now','-30 day') AND gravedad IN ('Grave','Muy grave','Alta')",(sid,))[0]['n']
    st=rows('SELECT curso_id FROM estudiantes WHERE id=?',(sid,)); vals=[]
    if st and st[0]['curso_id']:
        for a in rows('SELECT DISTINCT asignatura_id FROM curso_asignatura WHERE curso_id=? AND activa=1',(st[0]['curso_id'],)):
            v=_calificacion_final_base(sid,a['asignatura_id'])
            if v is not None: vals.append(v)
    prom=round(sum(vals)/len(vals),2) if vals else None
    venc=rows("SELECT COUNT(*) n FROM compromisos WHERE estudiante_id=? AND estado IN ('Pendiente','En progreso') AND fecha_limite IS NOT NULL AND fecha_limite<date('now')",(sid,))[0]['n']
    return {'asistencia':att,'registros_asistencia':total,'ausencias':aus,'incidencias_30_dias':inc,'incidencias_graves_30_dias':grave,'promedio':prom,'compromisos_vencidos':venc}

def _ensure_alerts(sid,u):
    if not can_access_student(u,sid): raise HTTPException(403,'No tienes acceso a este estudiante')
    st=rows('SELECT nombre,matricula FROM estudiantes WHERE id=?',(sid,))[0]; m=_risk_metrics(sid); rules=[]; um=float(_cfg('umbral_aprobacion','70'))
    if m['asistencia'] is not None and m['asistencia']<80: rules.append(('Asistencia','Alta' if m['asistencia']<70 else 'Media','Asistencia por debajo del umbral',m['asistencia'],80))
    if m['promedio'] is not None and m['promedio']<um: rules.append(('Rendimiento','Alta' if m['promedio']<60 else 'Media','Rendimiento académico bajo',m['promedio'],um))
    if m['incidencias_30_dias']>=3: rules.append(('Incidencias','Alta' if m['incidencias_30_dias']>=5 else 'Media','Incidencias recurrentes',m['incidencias_30_dias'],3))
    if m['incidencias_graves_30_dias']>=1: rules.append(('Incidencias graves','Alta','Incidencia grave reciente',m['incidencias_graves_30_dias'],1))
    if m['compromisos_vencidos']>=1: rules.append(('Seguimiento','Media','Compromisos vencidos',m['compromisos_vencidos'],1))
    c=conn(); out=[]
    for tipo,nivel,titulo,valor,threshold in rules:
        detail=f'Valor detectado: {valor}; umbral: {threshold}'
        old=c.execute("SELECT id FROM alertas_riesgo WHERE estudiante_id=? AND tipo=? AND titulo=? AND estado='Abierta'",(sid,tipo,titulo)).fetchone()
        if old: aid=old['id']
        else:
            cur=c.execute("INSERT INTO alertas_riesgo(estudiante_id,tipo,nivel,titulo,detalle,valor,umbral,estado,fecha_deteccion) VALUES(?,?,?,?,?,?,?,'Abierta',?)",(sid,tipo,nivel,titulo,detail,valor,threshold,now())); aid=cur.lastrowid
        out.append({'id':aid,'estudiante_id':sid,'estudiante':st['nombre'],'matricula':st['matricula'],'tipo':tipo,'nivel':nivel,'titulo':titulo,'detalle':detail,'valor':valor,'umbral':threshold,'estado':'Abierta'})
    c.commit(); c.close(); return out

@app.get('/api/alertas-avanzadas')
def alertas_avanzadas(estudiante_id:Optional[int]=None,estado:Optional[str]=None,u=Depends(current_user)):
    sts=_risk_students(u); sts=[x for x in sts if estudiante_id is None or x['id']==estudiante_id]
    if estudiante_id is not None and not sts: raise HTTPException(403,'No tienes acceso a este estudiante')
    for st in sts:_ensure_alerts(st['id'],u)
    ids=[x['id'] for x in sts]
    if not ids:return []
    ph=','.join('?'*len(ids)); q=f"SELECT a.*,e.nombre estudiante,e.matricula FROM alertas_riesgo a JOIN estudiantes e ON e.id=a.estudiante_id WHERE a.estudiante_id IN ({ph})"; args=ids
    if estado:q+=' AND a.estado=?';args=ids+[estado]
    q+=" ORDER BY CASE a.nivel WHEN 'Alta' THEN 1 WHEN 'Media' THEN 2 ELSE 3 END,a.fecha_deteccion DESC"
    return rows(q,args)

@app.put('/api/alertas-avanzadas/{alerta_id}')
def actualizar_alerta(alerta_id:int,x:AlertaEstadoIn,u=Depends(current_user)):
    r=rows('SELECT estudiante_id FROM alertas_riesgo WHERE id=?',(alerta_id,))
    if not r:raise HTTPException(404,'Alerta no encontrada')
    if not can_access_student(u,r[0]['estudiante_id']):raise HTTPException(403,'No tienes acceso a esta alerta')
    cierre=now() if x.estado.lower() in ('cerrada','resuelta','cerrado') else None
    c=conn();c.execute('UPDATE alertas_riesgo SET estado=?,responsable=?,observacion=?,fecha_cierre=? WHERE id=?',(x.estado,x.responsable,x.observacion,cierre,alerta_id));c.commit();c.close();return {'ok':True}

@app.get('/api/planes-intervencion')
def planes_intervencion(estudiante_id:Optional[int]=None,u=Depends(current_user)):
    sts=_risk_students(u);ids=[x['id'] for x in sts if estudiante_id is None or x['id']==estudiante_id]
    if estudiante_id is not None and not ids:raise HTTPException(403,'No tienes acceso a este estudiante')
    if not ids:return []
    ph=','.join('?'*len(ids));return rows(f'SELECT p.*,e.nombre estudiante,e.matricula FROM planes_intervencion p JOIN estudiantes e ON e.id=p.estudiante_id WHERE p.estudiante_id IN ({ph}) ORDER BY p.fecha_inicio DESC,p.id DESC',ids)

@app.post('/api/planes-intervencion')
def crear_plan(x:PlanIn,u=Depends(current_user)):
    if not can_access_student(u,x.estudiante_id):raise HTTPException(403,'No tienes acceso a este estudiante')
    if x.alerta_id and not rows('SELECT id FROM alertas_riesgo WHERE id=? AND estudiante_id=?',(x.alerta_id,x.estudiante_id)):raise HTTPException(400,'Alerta inválida')
    c=conn();cur=c.execute('INSERT INTO planes_intervencion(estudiante_id,alerta_id,titulo,objetivo,estrategia,responsable,fecha_inicio,fecha_revision,estado,resultado,observacion,fecha_registro) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',(x.estudiante_id,x.alerta_id,x.titulo,x.objetivo,x.estrategia,x.responsable,x.fecha_inicio.isoformat(),x.fecha_revision.isoformat() if x.fecha_revision else None,x.estado,x.resultado,x.observacion,now()));c.commit();i=cur.lastrowid;c.close();return {'id':i}

@app.put('/api/planes-intervencion/{plan_id}')
def actualizar_plan(plan_id:int,x:PlanIn,u=Depends(current_user)):
    r=rows('SELECT estudiante_id FROM planes_intervencion WHERE id=?',(plan_id,))
    if not r:raise HTTPException(404,'Plan no encontrado')
    if not can_access_student(u,r[0]['estudiante_id']):raise HTTPException(403,'No tienes acceso al plan')
    c=conn();c.execute('UPDATE planes_intervencion SET alerta_id=?,titulo=?,objetivo=?,estrategia=?,responsable=?,fecha_inicio=?,fecha_revision=?,estado=?,resultado=?,observacion=? WHERE id=?',(x.alerta_id,x.titulo,x.objetivo,x.estrategia,x.responsable,x.fecha_inicio.isoformat(),x.fecha_revision.isoformat() if x.fecha_revision else None,x.estado,x.resultado,x.observacion,plan_id));c.commit();c.close();return {'ok':True}

@app.get('/api/planes-intervencion/{plan_id}/acciones')
def acciones(plan_id:int,u=Depends(current_user)):
    r=rows('SELECT estudiante_id FROM planes_intervencion WHERE id=?',(plan_id,))
    if not r:raise HTTPException(404,'Plan no encontrado')
    if not can_access_student(u,r[0]['estudiante_id']):raise HTTPException(403,'No tienes acceso al plan')
    return rows('SELECT * FROM plan_acciones WHERE plan_id=? ORDER BY fecha_limite',(plan_id,))

@app.post('/api/planes-intervencion/acciones')
def crear_accion(x:PlanAccionIn,u=Depends(current_user)):
    r=rows('SELECT estudiante_id FROM planes_intervencion WHERE id=?',(x.plan_id,))
    if not r:raise HTTPException(404,'Plan no encontrado')
    if not can_access_student(u,r[0]['estudiante_id']):raise HTTPException(403,'No tienes acceso al plan')
    c=conn();cur=c.execute('INSERT INTO plan_acciones(plan_id,descripcion,responsable,fecha_limite,estado,evidencia,fecha_cumplimiento,observacion) VALUES(?,?,?,?,?,?,?,?)',(x.plan_id,x.descripcion,x.responsable,x.fecha_limite.isoformat() if x.fecha_limite else None,x.estado,x.evidencia,x.fecha_cumplimiento.isoformat() if x.fecha_cumplimiento else None,x.observacion));c.commit();i=cur.lastrowid;c.close();return {'id':i}

@app.put('/api/planes-intervencion/acciones/{accion_id}')
def actualizar_accion(accion_id:int,x:PlanAccionIn,u=Depends(current_user)):
    r=rows('SELECT p.estudiante_id FROM plan_acciones a JOIN planes_intervencion p ON p.id=a.plan_id WHERE a.id=?',(accion_id,))
    if not r:raise HTTPException(404,'Acción no encontrada')
    if not can_access_student(u,r[0]['estudiante_id']):raise HTTPException(403,'No tienes acceso a la acción')
    c=conn();c.execute('UPDATE plan_acciones SET descripcion=?,responsable=?,fecha_limite=?,estado=?,evidencia=?,fecha_cumplimiento=?,observacion=? WHERE id=?',(x.descripcion,x.responsable,x.fecha_limite.isoformat() if x.fecha_limite else None,x.estado,x.evidencia,x.fecha_cumplimiento.isoformat() if x.fecha_cumplimiento else None,x.observacion,accion_id));c.commit();c.close();return {'ok':True}

@app.get('/api/estudiante/{estudiante_id}/riesgo')
def riesgo(estudiante_id:int,u=Depends(current_user)):
    if not can_access_student(u,estudiante_id):raise HTTPException(403,'No tienes acceso a este estudiante')
    return {'metricas':_risk_metrics(estudiante_id),'alertas':_ensure_alerts(estudiante_id,u),'planes':planes_intervencion(estudiante_id,u)}
# ============================================================
# ETAPA 12 — COMUNICACIÓN INSTITUCIONAL Y NOTIFICACIONES
# ============================================================
def _comunicado_visible(c,u):
    if u['rol']=='admin': return True
    did=u.get('docente_id')
    if not did: return False
    return bool(rows("SELECT cd.id FROM comunicado_destinatarios cd LEFT JOIN cursos cu ON cu.id=cd.curso_id LEFT JOIN estudiantes e ON e.id=cd.estudiante_id WHERE cd.comunicado_id=? AND (cd.tipo_destinatario IN ('todos','docentes') OR (cd.tipo_destinatario='curso' AND cu.id IN (SELECT id FROM cursos WHERE docente_id=? UNION SELECT curso_id FROM curso_asignatura WHERE docente_id=? AND activa=1)) OR (cd.tipo_destinatario='estudiante' AND e.curso_id IN (SELECT id FROM cursos WHERE docente_id=? UNION SELECT curso_id FROM curso_asignatura WHERE docente_id=? AND activa=1))) LIMIT 1",(c['id'],did,did,did,did)))

@app.get('/api/comunicados')
def listar_comunicados(solo_no_leidos:bool=False,u=Depends(current_user)):
    base=rows("SELECT c.*,u.nombre autor_nombre,CASE WHEN cl.id IS NULL THEN 0 ELSE 1 END leido FROM comunicados c LEFT JOIN usuarios u ON u.id=c.autor_usuario_id LEFT JOIN comunicado_lecturas cl ON cl.comunicado_id=c.id AND cl.usuario_id=? WHERE c.activo=1 AND (c.fecha_expiracion IS NULL OR c.fecha_expiracion>=?) ORDER BY c.fecha_publicacion DESC",(u['id'],date.today().isoformat()))
    out=[x for x in base if _comunicado_visible(x,u)]
    return [x for x in out if not solo_no_leidos or not x['leido']]

@app.get('/api/comunicados/{comunicado_id}')
def detalle_comunicado(comunicado_id:int,u=Depends(current_user)):
    c=rows('SELECT c.*,u.nombre autor_nombre FROM comunicados c LEFT JOIN usuarios u ON u.id=c.autor_usuario_id WHERE c.id=?',(comunicado_id,))
    if not c: raise HTTPException(404,'Comunicado no encontrado')
    if not _comunicado_visible(c[0],u): raise HTTPException(403,'No tienes acceso a este comunicado')
    d=rows('SELECT cd.*,cu.nombre curso_nombre,e.nombre estudiante_nombre FROM comunicado_destinatarios cd LEFT JOIN cursos cu ON cu.id=cd.curso_id LEFT JOIN estudiantes e ON e.id=cd.estudiante_id WHERE cd.comunicado_id=?',(comunicado_id,))
    l=rows('SELECT cl.fecha_lectura,u.nombre FROM comunicado_lecturas cl JOIN usuarios u ON u.id=cl.usuario_id WHERE cl.comunicado_id=? ORDER BY cl.fecha_lectura DESC',(comunicado_id,)) if u['rol']=='admin' else []
    return {'comunicado':c[0],'destinatarios':d,'lecturas':l}

@app.post('/api/comunicados')
def crear_comunicado(x:ComunicadoIn,u=Depends(current_user)):
    if not x.destinatarios: raise HTTPException(400,'Debes indicar al menos un destinatario')
    for d in x.destinatarios:
        if d.tipo_destinatario not in ('todos','docentes','curso','estudiante'): raise HTTPException(400,'Tipo de destinatario no válido')
        if d.tipo_destinatario=='curso' and not d.curso_id: raise HTTPException(400,'Falta curso_id')
        if d.tipo_destinatario=='estudiante' and not d.estudiante_id: raise HTTPException(400,'Falta estudiante_id')
        if d.curso_id and not can_access_course(u,d.curso_id): raise HTTPException(403,'No tienes acceso a esa clase')
        if d.estudiante_id and not can_access_student(u,d.estudiante_id): raise HTTPException(403,'No tienes acceso a ese estudiante')
    c=conn();cur=c.execute('INSERT INTO comunicados(titulo,mensaje,tipo,prioridad,autor_usuario_id,fecha_publicacion,fecha_expiracion,activo) VALUES(?,?,?,?,?,?,?,?)',(x.titulo.strip(),x.mensaje,x.tipo,x.prioridad,u['id'],now(),x.fecha_expiracion.isoformat() if x.fecha_expiracion else None,1 if x.activo else 0));cid=cur.lastrowid
    for d in x.destinatarios:c.execute('INSERT OR IGNORE INTO comunicado_destinatarios(comunicado_id,tipo_destinatario,curso_id,estudiante_id) VALUES(?,?,?,?)',(cid,d.tipo_destinatario,d.curso_id,d.estudiante_id))
    c.commit();out=dict(c.execute('SELECT * FROM comunicados WHERE id=?',(cid,)).fetchone());c.close();return out

@app.post('/api/comunicados/{comunicado_id}/leer')
def marcar_comunicado_leido(comunicado_id:int,u=Depends(current_user)):
    c=rows('SELECT * FROM comunicados WHERE id=?',(comunicado_id,))
    if not c: raise HTTPException(404,'Comunicado no encontrado')
    if not _comunicado_visible(c[0],u): raise HTTPException(403,'No tienes acceso a este comunicado')
    db=conn();db.execute('INSERT OR IGNORE INTO comunicado_lecturas(comunicado_id,usuario_id,fecha_lectura) VALUES(?,?,?)',(comunicado_id,u['id'],now()));db.commit();db.close();return {'ok':True}

@app.delete('/api/comunicados/{comunicado_id}')
def desactivar_comunicado(comunicado_id:int,u=Depends(current_user)):
    c=rows('SELECT autor_usuario_id FROM comunicados WHERE id=?',(comunicado_id,))
    if not c: raise HTTPException(404,'Comunicado no encontrado')
    if u['rol']!='admin' and c[0]['autor_usuario_id']!=u['id']: raise HTTPException(403,'Solo el autor o un administrador puede desactivar el comunicado')
    db=conn();db.execute('UPDATE comunicados SET activo=0 WHERE id=?',(comunicado_id,));db.commit();db.close();return {'ok':True}

@app.get('/api/comunicados/resumen/no-leidos')
def resumen_no_leidos(u=Depends(current_user)):
    return {'cantidad':len(listar_comunicados(True,u))}


def _event_course(e):
    if e.get('curso_id'): return e['curso_id']
    if e.get('estudiante_id'):
        r=rows('SELECT curso_id FROM estudiantes WHERE id=?',(e['estudiante_id'],))
        return r[0]['curso_id'] if r else None
    return None

def _event_allowed(u,e):
    cid=_event_course(e)
    if cid is not None and not can_access_course(u,cid): return False
    if e.get('estudiante_id') is not None and not can_access_student(u,e['estudiante_id']): return False
    return True

@app.get('/api/agenda')
def agenda(desde:date, hasta:date, curso_id:Optional[int]=None, tipo:Optional[str]=None, u=Depends(current_user)):
    if hasta < desde: raise HTTPException(400,'El rango de fechas no es válido')
    if (hasta-desde).days>370: raise HTTPException(400,'El rango máximo es de 371 días')
    ids=allowed_course_ids(u); args=[desde.isoformat(),hasta.isoformat()]
    q='SELECT ev.*,c.nombre curso_nombre,c.grado,c.seccion,a.nombre asignatura_nombre,e.nombre estudiante_nombre,d.nombre creador_nombre FROM eventos_calendario ev LEFT JOIN cursos c ON c.id=ev.curso_id LEFT JOIN asignaturas a ON a.id=ev.asignatura_id LEFT JOIN estudiantes e ON e.id=ev.estudiante_id LEFT JOIN usuarios d ON d.id=ev.creado_por WHERE ev.fecha BETWEEN ? AND ?'
    if ids is not None:
        if not ids: return []
        q+=' AND (ev.curso_id IN ('+','.join('?'*len(ids))+') OR ev.curso_id IS NULL)'; args.extend(ids)
    if curso_id is not None:
        if not can_access_course(u,curso_id): raise HTTPException(403,'Sin acceso a la clase')
        q+=' AND (ev.curso_id=? OR ev.curso_id IS NULL)';args.append(curso_id)
    if tipo: q+=' AND ev.tipo=?';args.append(tipo)
    q+=' ORDER BY ev.fecha,ev.todo_el_dia DESC,ev.hora_inicio,ev.titulo COLLATE NOCASE'
    events=rows(q,args); extras=[]
    qa='SELECT ac.fecha,ac.titulo,ac.periodo,ac.estudiante_id,e.nombre estudiante_nombre,e.curso_id,a.nombre asignatura_nombre FROM actividades ac JOIN estudiantes e ON e.id=ac.estudiante_id LEFT JOIN asignaturas a ON a.id=ac.asignatura_id WHERE ac.fecha BETWEEN ? AND ?'; ar=[desde.isoformat(),hasta.isoformat()]
    if ids is not None:
        if ids: qa+=' AND e.curso_id IN ('+','.join('?'*len(ids))+')';ar.extend(ids)
        else: qa+=' AND 1=0'
    for x in rows(qa,ar): extras.append({'id':f"actividad-{x['fecha']}-{x['estudiante_id']}-{x['titulo']}",'origen':'actividad','titulo':x['titulo'],'tipo':'Evaluación','fecha':x['fecha'],'hora_inicio':None,'hora_fin':None,'todo_el_dia':1,'curso_id':x['curso_id'],'asignatura_id':None,'estudiante_id':x['estudiante_id'],'curso_nombre':None,'asignatura_nombre':x['asignatura_nombre'],'estudiante_nombre':x['estudiante_nombre'],'estado':'Registrada','solo_lectura':True,'periodo':x['periodo']})
    return [dict(x,solo_lectura=False,origen='evento') for x in events]+extras

@app.post('/api/agenda')
def crear_evento(x:EventoCalendarioIn,u=Depends(current_user)):
    data=x.model_dump(); cid=data.get('curso_id')
    if data.get('estudiante_id'):
        if not can_access_student(u,data['estudiante_id']): raise HTTPException(403,'Sin acceso al estudiante')
        if cid is None: cid=student_course(data['estudiante_id']); data['curso_id']=cid
    if cid is not None and not can_access_course(u,cid): raise HTTPException(403,'Sin acceso a la clase')
    if data.get('asignatura_id'):
        if not rows('SELECT id FROM asignaturas WHERE id=? AND activa=1',(data['asignatura_id'],)): raise HTTPException(404,'Asignatura no encontrada')
        if u['rol']!='admin' and not any(int(p['curso_id'])==int(cid) and int(p['asignatura_id'])==int(data['asignatura_id']) for p in (allowed_subject_pairs(u) or [])): raise HTTPException(403,'No tiene acceso a esa asignatura')
    c=conn();cur=c.execute('INSERT INTO eventos_calendario(titulo,tipo,fecha,hora_inicio,hora_fin,todo_el_dia,curso_id,asignatura_id,estudiante_id,ubicacion,descripcion,recordatorio_minutos,estado,creado_por,fecha_registro) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',(data['titulo'].strip(),data['tipo'],data['fecha'].isoformat(),data['hora_inicio'],data['hora_fin'],int(data['todo_el_dia']),cid,data['asignatura_id'],data['estudiante_id'],data['ubicacion'],data['descripcion'],data['recordatorio_minutos'],data['estado'],u['id'],now()));c.commit();r=dict(c.execute('SELECT * FROM eventos_calendario WHERE id=?',(cur.lastrowid,)).fetchone());c.close();return r

@app.put('/api/agenda/{evento_id}')
def actualizar_evento(evento_id:int,x:EventoCalendarioIn,u=Depends(current_user)):
    old=rows('SELECT * FROM eventos_calendario WHERE id=?',(evento_id,))
    if not old: raise HTTPException(404,'Evento no encontrado')
    if not _event_allowed(u,old[0]): raise HTTPException(403,'Sin acceso al evento')
    data=x.model_dump(); cid=data.get('curso_id')
    if data.get('estudiante_id'):
        if not can_access_student(u,data['estudiante_id']): raise HTTPException(403,'Sin acceso al estudiante')
        cid=cid or student_course(data['estudiante_id']);data['curso_id']=cid
    if cid is not None and not can_access_course(u,cid): raise HTTPException(403,'Sin acceso a la clase')
    c=conn();c.execute('UPDATE eventos_calendario SET titulo=?,tipo=?,fecha=?,hora_inicio=?,hora_fin=?,todo_el_dia=?,curso_id=?,asignatura_id=?,estudiante_id=?,ubicacion=?,descripcion=?,recordatorio_minutos=?,estado=? WHERE id=?',(data['titulo'].strip(),data['tipo'],data['fecha'].isoformat(),data['hora_inicio'],data['hora_fin'],int(data['todo_el_dia']),cid,data['asignatura_id'],data['estudiante_id'],data['ubicacion'],data['descripcion'],data['recordatorio_minutos'],data['estado'],evento_id));c.commit();r=dict(c.execute('SELECT * FROM eventos_calendario WHERE id=?',(evento_id,)).fetchone());c.close();return r

@app.delete('/api/agenda/{evento_id}')
def eliminar_evento(evento_id:int,u=Depends(current_user)):
    old=rows('SELECT * FROM eventos_calendario WHERE id=?',(evento_id,))
    if not old: raise HTTPException(404,'Evento no encontrado')
    if u['rol']!='admin' and old[0]['creado_por']!=u['id']: raise HTTPException(403,'Solo el creador o un administrador puede eliminar el evento')
    if not _event_allowed(u,old[0]): raise HTTPException(403,'Sin acceso al evento')
    c=conn();c.execute('DELETE FROM eventos_calendario WHERE id=?',(evento_id,));c.commit();c.close();return {'ok':True}

@app.get('/api/agenda/proximos')
def proximos_eventos(dias:int=7,u=Depends(current_user)):
    dias=max(1,min(dias,60)); ini=date.today(); fin=ini+timedelta(days=dias)
    return agenda(ini,fin,None,None,u)

def _plan_pair(u,curso_id,asignatura_id):
    if not can_access_course(u,curso_id): raise HTTPException(403,'Sin acceso a la clase')
    pair=rows('SELECT docente_id FROM curso_asignatura WHERE curso_id=? AND asignatura_id=? AND activa=1',(curso_id,asignatura_id))
    if not pair: raise HTTPException(400,'La asignatura no está asignada a la clase')
    if u['rol']!='admin' and pair[0]['docente_id']!=u.get('docente_id'): raise HTTPException(403,'Solo el docente responsable puede gestionar esta planificación')
    return pair[0]['docente_id']

@app.get('/api/planificaciones')
def listar_planificaciones(curso_id:Optional[int]=None,asignatura_id:Optional[int]=None,periodo:Optional[int]=None,anio_escolar_id:Optional[int]=None,u=Depends(current_user)):
    sql="""SELECT p.*,c.nombre curso_nombre,a.nombre asignatura_nombre,d.nombre docente_nombre,co.codigo competencia_codigo,co.nombre competencia_nombre
           FROM planificaciones p JOIN cursos c ON c.id=p.curso_id JOIN asignaturas a ON a.id=p.asignatura_id
           LEFT JOIN docentes d ON d.id=p.docente_id LEFT JOIN competencias co ON co.id=p.competencia_id WHERE 1=1"""; args=[]
    if u['rol']!='admin':
        ids=allowed_course_ids(u) or []
        if not ids: return []
        sql+=' AND p.curso_id IN ('+','.join('?'*len(ids))+')'; args+=ids
    for col,val in [('p.curso_id',curso_id),('p.asignatura_id',asignatura_id),('p.periodo',periodo),('p.anio_escolar_id',anio_escolar_id)]:
        if val is not None: sql+=' AND '+col+'=?'; args.append(val)
    sql+=' ORDER BY p.fecha_inicio IS NULL,p.fecha_inicio,p.periodo,p.id DESC'
    return rows(sql,tuple(args))

@app.post('/api/planificaciones')
def crear_planificacion(x:PlanificacionIn,u=Depends(current_user)):
    did=_plan_pair(u,x.curso_id,x.asignatura_id)
    if x.docente_id is not None and u['rol']!='admin' and x.docente_id!=u.get('docente_id'): raise HTTPException(403,'Docente no autorizado')
    if x.fecha_inicio and x.fecha_fin and x.fecha_fin<x.fecha_inicio: raise HTTPException(400,'La fecha final no puede ser anterior a la inicial')
    if x.competencia_id and not rows('SELECT id FROM competencias WHERE id=? AND asignatura_id=? AND activa=1',(x.competencia_id,x.asignatura_id)): raise HTTPException(400,'La competencia no pertenece a la asignatura')
    if not rows('SELECT id FROM anios_escolares WHERE id=?',(x.anio_escolar_id,)): raise HTTPException(404,'Año escolar no encontrado')
    did=x.docente_id or did; c=conn(); cur=c.cursor(); event_id=None
    if x.crear_evento and x.fecha_inicio:
        cur.execute("INSERT INTO eventos_calendario(titulo,tipo,fecha,todo_el_dia,curso_id,asignatura_id,descripcion,estado,creado_por,fecha_registro) VALUES(?,?,?,?,?,?,?,?,?,?)",(x.titulo,'Actividad',x.fecha_inicio.isoformat(),1,x.curso_id,x.asignatura_id,x.descripcion,'Programado',u['id'],now())); event_id=cur.lastrowid
    cur.execute("INSERT INTO planificaciones(curso_id,asignatura_id,docente_id,anio_escolar_id,periodo,competencia_id,titulo,fecha_inicio,fecha_fin,tipo_actividad,instrumento,descripcion,recursos,estado,evento_calendario_id,fecha_registro) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",(x.curso_id,x.asignatura_id,did,x.anio_escolar_id,x.periodo,x.competencia_id,x.titulo.strip(),x.fecha_inicio.isoformat() if x.fecha_inicio else None,x.fecha_fin.isoformat() if x.fecha_fin else None,x.tipo_actividad,x.instrumento,x.descripcion,x.recursos,x.estado,event_id,now())); pid=cur.lastrowid; c.commit(); out=dict(cur.execute('SELECT * FROM planificaciones WHERE id=?',(pid,)).fetchone()); c.close(); return out

@app.put('/api/planificaciones/{plan_id}')
def actualizar_planificacion(plan_id:int,x:PlanificacionIn,u=Depends(current_user)):
    old=rows('SELECT * FROM planificaciones WHERE id=?',(plan_id,))
    if not old: raise HTTPException(404,'Planificación no encontrada')
    _plan_pair(u,old[0]['curso_id'],old[0]['asignatura_id']); did=_plan_pair(u,x.curso_id,x.asignatura_id)
    if x.fecha_inicio and x.fecha_fin and x.fecha_fin<x.fecha_inicio: raise HTTPException(400,'La fecha final no puede ser anterior a la inicial')
    if x.competencia_id and not rows('SELECT id FROM competencias WHERE id=? AND asignatura_id=? AND activa=1',(x.competencia_id,x.asignatura_id)): raise HTTPException(400,'La competencia no pertenece a la asignatura')
    c=conn(); c.execute("UPDATE planificaciones SET curso_id=?,asignatura_id=?,docente_id=?,anio_escolar_id=?,periodo=?,competencia_id=?,titulo=?,fecha_inicio=?,fecha_fin=?,tipo_actividad=?,instrumento=?,descripcion=?,recursos=?,estado=? WHERE id=?",(x.curso_id,x.asignatura_id,x.docente_id or did,x.anio_escolar_id,x.periodo,x.competencia_id,x.titulo.strip(),x.fecha_inicio.isoformat() if x.fecha_inicio else None,x.fecha_fin.isoformat() if x.fecha_fin else None,x.tipo_actividad,x.instrumento,x.descripcion,x.recursos,x.estado,plan_id)); c.commit(); out=dict(c.execute('SELECT * FROM planificaciones WHERE id=?',(plan_id,)).fetchone()); c.close(); return out

@app.delete('/api/planificaciones/{plan_id}')
def eliminar_planificacion(plan_id:int,u=Depends(current_user)):
    old=rows('SELECT * FROM planificaciones WHERE id=?',(plan_id,))
    if not old: raise HTTPException(404,'Planificación no encontrada')
    _plan_pair(u,old[0]['curso_id'],old[0]['asignatura_id']); c=conn(); c.execute('DELETE FROM planificaciones WHERE id=?',(plan_id,)); c.commit(); c.close(); return {'ok':True}

@app.get('/api/horario')
def listar_horario(curso_id:Optional[int]=None, docente_id:Optional[int]=None, u=Depends(current_user)):
    params=[]; where=['h.activa=1']
    if u['rol']!='admin':
        allowed=allowed_course_ids(u) or []
        if not allowed: return []
        where.append('h.curso_id IN ('+','.join('?' for _ in allowed)+')'); params.extend(allowed)
    if curso_id is not None:
        if not can_access_course(u,curso_id): raise HTTPException(403,'Sin acceso a la clase')
        where.append('h.curso_id=?');params.append(curso_id)
    if docente_id is not None:
        if u['rol']!='admin' and docente_id!=u.get('docente_id'): raise HTTPException(403,'Sin acceso al docente')
        where.append('h.docente_id=?');params.append(docente_id)
    return rows("SELECT h.*,c.nombre curso_nombre,c.grado,c.seccion,a.nombre asignatura_nombre,d.nombre docente_nombre FROM horario_escolar h JOIN cursos c ON c.id=h.curso_id JOIN asignaturas a ON a.id=h.asignatura_id LEFT JOIN docentes d ON d.id=h.docente_id WHERE "+' AND '.join(where)+" ORDER BY h.dia_semana,h.hora_inicio,c.nombre",tuple(params))

@app.post('/api/horario')
def crear_horario(x:HorarioIn,u=Depends(current_user)):
    if not can_access_course(u,x.curso_id): raise HTTPException(403,'Sin acceso a la clase')
    pair=rows('SELECT docente_id FROM curso_asignatura WHERE curso_id=? AND asignatura_id=? AND activa=1',(x.curso_id,x.asignatura_id))
    if not pair: raise HTTPException(400,'La asignatura no está asignada a la clase')
    did=x.docente_id or pair[0]['docente_id']
    if u['rol']!='admin' and did!=u.get('docente_id'): raise HTTPException(403,'Solo puedes programar tu propia asignatura')
    if x.hora_fin<=x.hora_inicio: raise HTTPException(400,'La hora final debe ser posterior a la inicial')
    conflict=rows("SELECT h.id,c.nombre curso_nombre,a.nombre asignatura_nombre,d.nombre docente_nombre FROM horario_escolar h JOIN cursos c ON c.id=h.curso_id JOIN asignaturas a ON a.id=h.asignatura_id LEFT JOIN docentes d ON d.id=h.docente_id WHERE h.activa=1 AND h.dia_semana=? AND h.hora_inicio < ? AND h.hora_fin > ? AND (h.curso_id=? OR h.docente_id=?)",(x.dia_semana,x.hora_fin,x.hora_inicio,x.curso_id,did))
    if conflict: raise HTTPException(409,detail={'message':'Conflicto de horario','conflictos':conflict})
    c=conn();cur=c.execute('INSERT INTO horario_escolar(dia_semana,hora_inicio,hora_fin,curso_id,asignatura_id,docente_id,aula,bloque,activa,fecha_registro) VALUES(?,?,?,?,?,?,?,?,?,?)',(x.dia_semana,x.hora_inicio,x.hora_fin,x.curso_id,x.asignatura_id,did,x.aula,x.bloque,int(x.activa),now()));c.commit();r=dict(c.execute('SELECT * FROM horario_escolar WHERE id=?',(cur.lastrowid,)).fetchone());c.close();return r

@app.put('/api/horario/{horario_id}')
def actualizar_horario(horario_id:int,x:HorarioIn,u=Depends(current_user)):
    old=rows('SELECT * FROM horario_escolar WHERE id=?',(horario_id,))
    if not old: raise HTTPException(404,'Horario no encontrado')
    if not can_access_course(u,old[0]['curso_id']) or not can_access_course(u,x.curso_id): raise HTTPException(403,'Sin acceso al horario')
    if u['rol']!='admin' and old[0]['docente_id']!=u.get('docente_id'): raise HTTPException(403,'Solo el docente responsable puede editar este horario')
    pair=rows('SELECT docente_id FROM curso_asignatura WHERE curso_id=? AND asignatura_id=? AND activa=1',(x.curso_id,x.asignatura_id))
    if not pair: raise HTTPException(400,'La asignatura no está asignada a la clase')
    did=x.docente_id or pair[0]['docente_id']
    if x.hora_fin<=x.hora_inicio: raise HTTPException(400,'La hora final debe ser posterior a la inicial')
    conflict=rows('SELECT h.id FROM horario_escolar h WHERE h.activa=1 AND h.id<>? AND h.dia_semana=? AND h.hora_inicio < ? AND h.hora_fin > ? AND (h.curso_id=? OR h.docente_id=?)',(horario_id,x.dia_semana,x.hora_fin,x.hora_inicio,x.curso_id,did))
    if conflict: raise HTTPException(409,'Conflicto de horario')
    c=conn();c.execute('UPDATE horario_escolar SET dia_semana=?,hora_inicio=?,hora_fin=?,curso_id=?,asignatura_id=?,docente_id=?,aula=?,bloque=?,activa=? WHERE id=?',(x.dia_semana,x.hora_inicio,x.hora_fin,x.curso_id,x.asignatura_id,did,x.aula,x.bloque,int(x.activa),horario_id));c.commit();r=dict(c.execute('SELECT * FROM horario_escolar WHERE id=?',(horario_id,)).fetchone());c.close();return r

@app.delete('/api/horario/{horario_id}')
def eliminar_horario(horario_id:int,u=Depends(current_user)):
    old=rows('SELECT * FROM horario_escolar WHERE id=?',(horario_id,))
    if not old: raise HTTPException(404,'Horario no encontrado')
    if not can_access_course(u,old[0]['curso_id']): raise HTTPException(403,'Sin acceso al horario')
    if u['rol']!='admin' and old[0]['docente_id']!=u.get('docente_id'): raise HTTPException(403,'Solo el docente responsable puede eliminar este horario')
    c=conn();c.execute('UPDATE horario_escolar SET activa=0 WHERE id=?',(horario_id,));c.commit();c.close();return {'ok':True}

@app.get('/api/sistema/estado')
def sistema_estado(u=Depends(current_user)):
    c=conn(); integrity=c.execute('PRAGMA integrity_check').fetchone()[0]; fk_errors=len(c.execute('PRAGMA foreign_key_check').fetchall())
    tables=['usuarios','docentes','cursos','asignaturas','estudiantes','actividades','calificaciones','asistencia','incidencias','permisos','eventos_calendario','horario_escolar','planificaciones','auditoria']
    counts={}
    for t in tables:
        try: counts[t]=c.execute(f'SELECT COUNT(*) FROM {t}').fetchone()[0]
        except Exception: counts[t]=None
    c.close()
    return {'version':app.version,'base_datos':{'integridad':integrity,'errores_claves_foraneas':fk_errors,'ruta':str(DB_PATH)},'conteos':counts,'rol':u['rol']}

@app.get('/api/db/info')
def db_info(u=Depends(current_user)):
    c=conn();ts=c.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name").fetchall();counts={t[0]:c.execute(f'SELECT COUNT(*) FROM "{t[0]}"').fetchone()[0] for t in ts};c.close();return {'path':str(DB_PATH),'tables':counts}


# Frontend autocontenido: se sirve desde FastAPI; no requiere Node.js/Vite.
STATIC_DIR=BASE_DIR/'frontend_dist'
if STATIC_DIR.exists():
    app.mount('/assets', StaticFiles(directory=STATIC_DIR/'assets'), name='assets') if (STATIC_DIR/'assets').exists() else None
    @app.get('/', include_in_schema=False)
    def frontend_index():
        return FileResponse(STATIC_DIR/'index.html')
    @app.get('/{full_path:path}', include_in_schema=False)
    def frontend_fallback(full_path: str):
        # Las rutas /api/* ya fueron registradas arriba; el resto cae al frontend.
        target=STATIC_DIR/full_path
        if target.is_file(): return FileResponse(target)
        return FileResponse(STATIC_DIR/'index.html')
